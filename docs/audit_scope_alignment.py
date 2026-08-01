#!/usr/bin/env python3
"""Audit the canonical 56-use-case catalogue across retake artifacts."""

from __future__ import annotations

import json
import re
import sys
from pathlib import Path

from docx import Document

from update_imported_rds_sds import CLASS_IMAGE_NAMES, FINAL_USE_CASES


ROOT = Path(__file__).resolve().parents[1]
EXPECTED = [(number, use_case[4]) for number, use_case in enumerate(FINAL_USE_CASES, start=1)]


def fail(message: str, failures: list[str]) -> None:
    failures.append(message)


def backlog_rows(path: Path):
    for line in path.read_text(encoding="utf-8").splitlines():
        record = json.loads(line)
        if record.get("kind") == "table" and record.get("sheet") == "Backlog" and record.get("address") == "A1:J57":
            return [(int(row[0].split("-")[1]), row[4]) for row in record["values"][1:]]
    raise RuntimeError(f"Backlog table A1:J57 not found in {path}")


def rds_rows(path: Path):
    rows = []
    for paragraph in Document(path).paragraphs:
        match = re.fullmatch(r"UC-(\d{2})\.\s+(.+)", " ".join(paragraph.text.split()))
        if match:
            rows.append((int(match.group(1)), match.group(2)))
    return rows


def sds_rows(path: Path):
    rows = []
    in_designs = False
    for paragraph in Document(path).paragraphs:
        text = " ".join(paragraph.text.split())
        if text == "II. Code Designs":
            in_designs = True
            continue
        if not in_designs or not paragraph.style or paragraph.style.name != "Heading 2":
            continue
        match = re.fullmatch(r"(\d+)\.\s+(.+)", text)
        if match:
            rows.append((int(match.group(1)), match.group(2)))
    return rows


def sds_query_titles(path: Path):
    titles = []
    for paragraph in Document(path).paragraphs:
        match = re.fullmatch(r"== (.+) ==", " ".join(paragraph.text.split()))
        if match:
            titles.append(match.group(1))
    return [(number, title) for number, title in enumerate(titles, start=1)]


def sequence_rows(directory: Path):
    rows = []
    for path in sorted(directory.glob("UC-*.puml")):
        text = path.read_text(encoding="utf-8")
        match = re.search(r"^== UC-(\d{2}): (.+) ==$", text, flags=re.MULTILINE)
        if not match:
            raise RuntimeError(f"Missing canonical title in {path}")
        rows.append((int(match.group(1)), match.group(2)))
    return rows


def compare(label: str, actual, failures: list[str]) -> None:
    if actual == EXPECTED:
        print(f"{label}: 56 IDs/names aligned")
        return
    for index in range(max(len(actual), len(EXPECTED))):
        expected = EXPECTED[index] if index < len(EXPECTED) else None
        received = actual[index] if index < len(actual) else None
        if expected != received:
            fail(f"{label} row {index + 1}: expected {expected!r}, received {received!r}", failures)


def main() -> int:
    failures: list[str] = []
    compare(
        "Backlog",
        backlog_rows(ROOT / "outputs/retake-refactor/GoalZone_Backlog_Retake_Aligned_2026-08-01.xlsx.inspect.ndjson"),
        failures,
    )
    compare("RDS", rds_rows(ROOT / "output/doc/GoalZone_RDS_Retake_Aligned_2026-08-01.docx"), failures)
    compare("SDS", sds_rows(ROOT / "output/doc/GoalZone_SDS_Retake_Aligned_2026-08-01.docx"), failures)
    compare("SDS query sections", sds_query_titles(ROOT / "output/doc/GoalZone_SDS_Retake_Aligned_2026-08-01.docx"), failures)
    compare("Sequence diagrams", sequence_rows(ROOT / "docs/Sequence"), failures)

    expected_class_sources = {Path(name).with_suffix(".puml").name for name in CLASS_IMAGE_NAMES}
    actual_class_sources = {path.name for path in (ROOT / "docs/Class Diagram").glob("*.puml")}
    if expected_class_sources == actual_class_sources and len(actual_class_sources) == 56:
        print("Class diagrams: 56 mapped sources aligned")
    else:
        fail(f"Class diagram set differs: missing={sorted(expected_class_sources - actual_class_sources)}, extra={sorted(actual_class_sources - expected_class_sources)}", failures)

    if failures:
        print("Scope alignment audit: FAIL")
        for item in failures:
            print(f"- {item}")
        return 1
    print("Scope alignment audit: PASS")
    return 0


if __name__ == "__main__":
    sys.exit(main())
