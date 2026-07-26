from __future__ import annotations

import argparse
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parents[1]
RDS_PATH = ROOT / "output" / "doc" / "RDS Document_G03_Final_Code_First_Aligned.docx"
SDS_PATH = ROOT / "output" / "doc" / "SDS Document_G03_Final_Code_First_Aligned.docx"


PRIMARY_ACTOR_OVERRIDES = {
    "UC-04": "Customer, Staff, or Admin",
    "UC-09": "Admin",
    "UC-15": "Admin or Staff",
    "UC-16": "Staff or Admin",
    "UC-18": "Customer or Staff",
    "UC-25": "Owning Customer, Staff, or Admin",
    "UC-28": "Customer or Staff",
    "UC-37": "Guest, Customer, or Staff",
    "UC-38": "Customer or Staff",
    "UC-40": "Customer",
    "UC-42": "Customer or Time",
    "UC-43": "Customer, Staff, or Admin",
    "UC-44": "Customer or Staff",
    "UC-45": "Customer, Staff, or Admin",
    "UC-50": "Guest or Customer",
    "UC-52": "Customer or Staff",
    "UC-55": "Guest or Customer",
    "UC-56": "Customer or Staff",
    "UC-57": "Time",
    "UC-58": "Customer or Staff",
    "UC-62": "Customer",
    "UC-63": "Guest or Customer",
}


SECONDARY_ACTORS = {
    "UC-01": "Email service",
    "UC-02": "None",
    "UC-03": "None",
    "UC-04": "None",
    "UC-05": "None",
    "UC-06": "Email service",
    "UC-07": "None",
    "UC-08": "Email service",
    "UC-09": "None",
    "UC-10": "None",
    "UC-11": "None",
    "UC-12": "None",
    "UC-13": "None",
    "UC-14": "None",
    "UC-15": "None",
    "UC-16": "None",
    "UC-17": "None",
    "UC-18": "None",
    "UC-19": "None",
    "UC-20": "None",
    "UC-21": "None",
    "UC-22": "Customer",
    "UC-23": "PayPal Sandbox",
    "UC-24": "Customer",
    "UC-25": "None",
    "UC-26": "None",
    "UC-27": "None",
    "UC-28": "None",
    "UC-29": "Customer",
    "UC-30": "Customer when Staff reschedules on the Customer's behalf",
    "UC-31": "None",
    "UC-32": "Customer when Staff cancels on the Customer's behalf",
    "UC-33": "None",
    "UC-34": "Customer",
    "UC-35": "None",
    "UC-36": "Customer",
    "UC-37": "None",
    "UC-38": "None",
    "UC-39": "PayPal Sandbox",
    "UC-40": "PayPal Sandbox",
    "UC-41": "None",
    "UC-42": "PayPal Sandbox",
    "UC-43": "None",
    "UC-44": "None",
    "UC-45": "None",
    "UC-46": "PayPal Sandbox",
    "UC-47": "Customer; PayPal Sandbox",
    "UC-48": "None",
    "UC-49": "None",
    "UC-50": "None",
    "UC-51": "None",
    "UC-52": "None",
    "UC-53": "None",
    "UC-54": "None",
    "UC-55": "None",
    "UC-56": "None",
    "UC-57": "Customer",
    "UC-58": "None",
    "UC-59": "None",
    "UC-60": "None",
    "UC-61": "None",
    "UC-62": "Gemini API",
    "UC-63": "None",
}


RDS_CONTENT_OVERRIDES = {
    "UC-09": {
        "Trigger:": "Admin selects a Customer from People & access to review operational history.",
        "Assumptions:": "Recent-list limits are sufficient for the Admin workflow.",
    },
    "UC-62": {
        "Trigger:": "An authenticated Customer opens /assistant and submits a natural-language availability request.",
        "Description:": (
            "Gemini extracts date, preferred time, field type, and budget from an English question, "
            "while GoalZone's deterministic slot service remains the source of truth for availability, "
            "prices, and eligible promotions."
        ),
        "Preconditions:": (
            "PRE-1: The Customer is authenticated.\n"
            "PRE-2: The Gemini API key and model are configured.\n"
            "PRE-3: Active field, pricing, and slot-generation data exist."
        ),
        "Postconditions:": (
            "POST-1: A grounded answer, interpreted criteria, and up to eight live ranked slots are returned.\n"
            "POST-2: No Booking or Payment is created by the assistant request."
        ),
        "Normal Flow:": (
            "A. Ask the Gemini availability assistant\n"
            "1. The Customer asks for a field in plain English.\n"
            "2. The client posts the question to /api/assistant/availability.\n"
            "3. The server verifies the authenticated Customer and asks Gemini to extract constrained search criteria.\n"
            "4. The server validates the date, time, field type, and budget.\n"
            "5. FieldOperationService materializes and ranks only live available slots.\n"
            "6. Gemini phrases a short answer using only the verified criteria and slots.\n"
            "7. The client displays the grounded answer and booking-ready recommendations."
        ),
        "Alternative Flow:": (
            "A1: The visitor uses the deterministic filters\n"
            "1. Guest or Customer selects date, time, field type, and budget manually.\n"
            "2. The client calls GET /api/slots/suggestions without requiring Gemini.\n"
            "A2: No matching slot exists\n"
            "1. The assistant reports no result and suggests broadening one criterion.\n"
            "A3: Optional criteria are omitted\n"
            "1. The date defaults to tomorrow and omitted filters remain unrestricted."
        ),
        "Exceptions:": (
            "EXC-01: Missing Gemini configuration returns service unavailable.\n"
            "EXC-02: Gemini rejection, rate limit, timeout, or invalid output returns a gateway error without fabricated slots.\n"
            "EXC-03: A past date or invalid request returns bad request."
        ),
        "Other Information:": (
            "Gemini interprets and phrases the request only. GoalZone data remains authoritative for every "
            "field, slot, price, and promotion shown. UC-63 handles selection and booking handoff."
        ),
        "Assumptions:": (
            "The Customer asks in English using at most 500 characters, and the configured Gemini service is reachable."
        ),
    },
    "UC-63": {
        "Trigger:": "A Guest or Customer selects Book this slot on a ranked recommendation.",
        "Preconditions:": (
            "PRE-1: UC-62 or the deterministic filter returned at least one live available suggestion.\n"
            "PRE-2: The visitor selects a result."
        ),
        "Normal Flow:": (
            "A. Continue a suggested slot into booking\n"
            "1. The client stores the suggested slot id and date in sessionStorage.\n"
            "2. The client navigates to /booking.\n"
            "3. The Booking page reloads live slots for the suggested date.\n"
            "4. If still available, the app selects the exact suggested slot and clears the handoff value.\n"
            "5. A Guest signs in or registers when required; the Customer then continues the normal checkout flow."
        ),
    },
}


def clean(value: str) -> str:
    return " ".join(value.split())


def set_cell_text_preserving_format(cell, value: str) -> None:
    paragraph = cell.paragraphs[0]
    runs = paragraph.runs
    if runs:
        runs[0].text = value
        for run in runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(value)
    for extra_paragraph in cell.paragraphs[1:]:
        for run in extra_paragraph.runs:
            run.text = ""


def update_rds_actors(path: Path) -> None:
    document = Document(path)
    updated: set[str] = set()
    for table in document.tables:
        uc_id = None
        actor_row = None
        for row in table.rows[:8]:
            values = [clean(cell.text) for cell in row.cells]
            if values and values[0].startswith("UC ID and Name:") and len(values) > 1:
                uc_id = values[1].split(maxsplit=1)[0]
            if values and values[0].startswith("Primary Actor:"):
                actor_row = row
        if not uc_id or actor_row is None or uc_id not in SECONDARY_ACTORS:
            continue
        if uc_id in PRIMARY_ACTOR_OVERRIDES:
            set_cell_text_preserving_format(actor_row.cells[1], PRIMARY_ACTOR_OVERRIDES[uc_id])
        set_cell_text_preserving_format(actor_row.cells[3], SECONDARY_ACTORS[uc_id])
        for row in table.rows:
            label = clean(row.cells[0].text)
            replacement = RDS_CONTENT_OVERRIDES.get(uc_id, {}).get(label)
            if replacement is not None:
                set_cell_text_preserving_format(row.cells[1], replacement)
        updated.add(uc_id)
    missing = sorted(set(SECONDARY_ACTORS) - updated)
    if missing:
        raise RuntimeError(f"RDS actor rows not found: {', '.join(missing)}")
    document.save(path)


def ensure_sds_toc_styles(sds: Document) -> tuple[str, str]:
    style_ids: list[str] = []
    for level in (1, 2):
        name = f"TOC {level}"
        try:
            style = sds.styles[name]
        except KeyError:
            style = sds.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = sds.styles["Normal"]
        style.font.name = "Times New Roman"
        style.font.size = Pt(12)
        style.font.italic = True
        formatting = style.paragraph_format
        formatting.left_indent = Inches(0 if level == 1 else 0.25)
        formatting.first_line_indent = Inches(0)
        formatting.space_before = Pt(0)
        formatting.space_after = Pt(0)
        formatting.line_spacing = 1
        formatting.tab_stops.clear_all()
        formatting.tab_stops.add_tab_stop(
            Inches(6.5),
            WD_TAB_ALIGNMENT.RIGHT,
            WD_TAB_LEADER.DOTS,
        )
        style_ids.append(style.style_id)
    return style_ids[0], style_ids[1]


def replace_title_with_rds_format(rds: Document, sds: Document) -> None:
    rds_title = next(paragraph for paragraph in rds.paragraphs if clean(paragraph.text) == "Contents")
    sds_title = next(
        paragraph
        for paragraph in sds.paragraphs
        if clean(paragraph.text) in {"Table of Contents", "Contents"}
    )
    sds_title._p.getparent().replace(sds_title._p, deepcopy(rds_title._p))


def find_toc_sdt(document: Document):
    for sdt in document.element.body.iter(qn("w:sdt")):
        instruction = "".join(
            element.text or "" for element in sdt.iter(qn("w:instrText"))
        )
        if "TOC" in instruction:
            return sdt
    raise RuntimeError("SDS table of contents content control not found")


def page_map_from_pdf(pdf_path: Path, headings: list[str], skip_pages: int) -> dict[str, int]:
    reader = PdfReader(str(pdf_path))
    pages = [(page.extract_text() or "") for page in reader.pages]
    result: dict[str, int] = {}
    for heading in headings:
        for page_number, text in enumerate(pages[skip_pages:], start=skip_pages + 1):
            if heading in text:
                result[heading] = page_number
                break
        if heading not in result:
            raise RuntimeError(f"Heading not found in rendered SDS PDF: {heading}")
    return result


def ordered_heading_pages_from_pdf(
    pdf_path: Path,
    headings: list[str],
    ignored_pages: set[int],
) -> list[int]:
    reader = PdfReader(str(pdf_path))
    pages = [" ".join((page.extract_text() or "").split()) for page in reader.pages]
    result: list[int] = []
    page_cursor = 1
    for heading in headings:
        normalized_heading = clean(heading)
        found = None
        for page_number in range(page_cursor, len(pages) + 1):
            if page_number in ignored_pages:
                continue
            if normalized_heading in pages[page_number - 1]:
                found = page_number
                break
        if found is None:
            raise RuntimeError(f"RDS heading not found in rendered PDF: {heading}")
        result.append(found)
        page_cursor = found
    return result


def refresh_rds_cached_toc_pages(path: Path, pdf_path: Path) -> None:
    document = Document(path)
    heading_titles = [
        clean(paragraph.text)
        for paragraph in document.paragraphs
        if paragraph.style.name in {"Heading 1", "Heading 2", "Heading 3"}
        and clean(paragraph.text)
    ]
    heading_pages = ordered_heading_pages_from_pdf(
        pdf_path,
        heading_titles,
        ignored_pages={3},
    )
    toc = find_toc_sdt(document)
    heading_cursor = 0
    updated = 0
    for paragraph in toc.iter(qn("w:p")):
        text_nodes = list(paragraph.iter(qn("w:t")))
        visible = "".join(node.text or "" for node in text_nodes).strip()
        if not visible:
            continue
        page_node = next(
            (node for node in reversed(text_nodes) if (node.text or "").strip().isdigit()),
            None,
        )
        if page_node is None:
            continue
        page_text = (page_node.text or "").strip()
        title = visible[: -len(page_text)].strip()
        match_index = None
        for index in range(heading_cursor, len(heading_titles)):
            if heading_titles[index] == title:
                match_index = index
                break
        if match_index is None:
            raise RuntimeError(f"RDS TOC heading not matched: {title}")
        page_node.text = str(heading_pages[match_index])
        heading_cursor = match_index + 1
        updated += 1
    if updated == 0:
        raise RuntimeError("No RDS TOC page values were refreshed")
    document.save(path)


def add_text_run(paragraph, text: str, italic: bool = True) -> None:
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    if italic:
        properties.append(OxmlElement("w:i"))
        properties.append(OxmlElement("w:iCs"))
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Times New Roman")
    fonts.set(qn("w:hAnsi"), "Times New Roman")
    fonts.set(qn("w:cs"), "Times New Roman")
    properties.append(fonts)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "24")
    properties.append(size)
    complex_size = OxmlElement("w:szCs")
    complex_size.set(qn("w:val"), "24")
    properties.append(complex_size)
    run.append(properties)
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    paragraph.append(run)


def add_tab_run(paragraph) -> None:
    run = OxmlElement("w:r")
    tab = OxmlElement("w:tab")
    run.append(tab)
    paragraph.append(run)


def add_field_character(paragraph, kind: str, dirty: bool = False) -> None:
    run = OxmlElement("w:r")
    field = OxmlElement("w:fldChar")
    field.set(qn("w:fldCharType"), kind)
    if dirty:
        field.set(qn("w:dirty"), "true")
    run.append(field)
    paragraph.append(run)


def add_instruction(paragraph, instruction: str) -> None:
    run = OxmlElement("w:r")
    node = OxmlElement("w:instrText")
    node.set(qn("xml:space"), "preserve")
    node.text = instruction
    run.append(node)
    paragraph.append(run)


def toc_entry(
    title: str,
    page_number: int,
    style_id: str,
    level: int,
    first: bool,
    last: bool,
):
    paragraph = OxmlElement("w:p")
    properties = OxmlElement("w:pPr")
    style = OxmlElement("w:pStyle")
    style.set(qn("w:val"), style_id)
    properties.append(style)
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "right")
    tab.set(qn("w:leader"), "dot")
    tab.set(qn("w:pos"), "9360")
    tabs.append(tab)
    properties.append(tabs)
    indentation = OxmlElement("w:ind")
    indentation.set(qn("w:left"), "0" if level == 1 else "360")
    properties.append(indentation)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:before"), "0")
    spacing.set(qn("w:after"), "0")
    spacing.set(qn("w:line"), "240")
    spacing.set(qn("w:lineRule"), "auto")
    properties.append(spacing)
    paragraph.append(properties)
    if first:
        add_field_character(paragraph, "begin", dirty=True)
        add_instruction(paragraph, ' TOC \\o "1-2" \\h \\z \\u ')
        add_field_character(paragraph, "separate")
    add_text_run(paragraph, title)
    add_tab_run(paragraph)
    add_text_run(paragraph, str(page_number))
    if last:
        add_field_character(paragraph, "end")
    return paragraph


def update_sds_toc(
    path: Path,
    rds_path: Path,
    pdf_path: Path,
    page_offset: int,
    pdf_skip_pages: int,
) -> None:
    rds = Document(rds_path)
    sds = Document(path)
    replace_title_with_rds_format(rds, sds)
    toc_style_ids = ensure_sds_toc_styles(sds)

    headings: list[tuple[str, int]] = []
    for paragraph in sds.paragraphs:
        title = clean(paragraph.text)
        if not title:
            continue
        if paragraph.style.name == "Heading 1":
            headings.append((title, 1))
        elif paragraph.style.name == "Heading 2":
            headings.append((title, 2))
    page_map = page_map_from_pdf(
        pdf_path,
        [title for title, _ in headings],
        skip_pages=pdf_skip_pages,
    )

    sdt = find_toc_sdt(sds)
    content = sdt.find(qn("w:sdtContent"))
    for child in list(content):
        content.remove(child)
    for index, (title, level) in enumerate(headings):
        content.append(
            toc_entry(
                title=title,
                page_number=page_map[title] + page_offset,
                style_id=toc_style_ids[level - 1],
                level=level,
                first=index == 0,
                last=index == len(headings) - 1,
            )
        )
    settings = sds.settings.element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")
    sds.save(path)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--sds-pdf", type=Path)
    parser.add_argument("--rds-pdf", type=Path)
    parser.add_argument("--toc-page-offset", type=int, default=0)
    parser.add_argument("--pdf-skip-pages", type=int, default=3)
    parser.add_argument("--skip-rds-actors", action="store_true")
    parser.add_argument("--skip-sds-toc", action="store_true")
    parser.add_argument("--refresh-rds-toc-pages", action="store_true")
    args = parser.parse_args()
    if not args.skip_rds_actors:
        update_rds_actors(RDS_PATH)
    if args.refresh_rds_toc_pages:
        if args.rds_pdf is None:
            parser.error("--rds-pdf is required with --refresh-rds-toc-pages")
        refresh_rds_cached_toc_pages(RDS_PATH, args.rds_pdf)
    if not args.skip_sds_toc:
        if args.sds_pdf is None:
            parser.error("--sds-pdf is required unless --skip-sds-toc is used")
        update_sds_toc(
            SDS_PATH,
            RDS_PATH,
            args.sds_pdf,
            args.toc_page_offset,
            args.pdf_skip_pages,
        )
    print(RDS_PATH)
    print(SDS_PATH)


if __name__ == "__main__":
    main()
