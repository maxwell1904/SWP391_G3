#!/usr/bin/env python3
"""Insert the code-first UC class and sequence diagrams into the SDS DOCX."""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches
from docx.text.paragraph import Paragraph
from PIL import Image


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
CLASS_DIR = DOCS / "Class Diagram"
SEQUENCE_DIR = DOCS / "Sequence"


CLASS_FILES = {
    1: "register.puml",
    2: "login.puml",
    3: "logout.puml",
    4: "manage-personal-profile.puml",
    5: "change-password.puml",
    6: "forgot-reset-password.puml",
    7: "manage-customer-accounts.puml",
    8: "manage-staff-accounts.puml",
    9: "view-customer-activity-status.puml",
    10: "restrict-unrestrict-customer-booking-ability.puml",
    11: "view-field-list.puml",
    12: "view-field-detail.puml",
    13: "search-available-fields.puml",
    14: "manage-football-fields.puml",
    15: "manage-field-pricing-by-time-range.puml",
    16: "manage-unavailable-slots.puml",
    17: "view-field-operation-calendar.puml",
    18: "manage-extra-services.puml",
    19: "check-extra-service-availability.puml",
    20: "add-extra-services-to-booking.puml",
    21: "update-extra-services-before-check-in.puml",
    22: "report-field-service-issue.puml",
    23: "resolve-field-service-issue.puml",
    24: "create-online-booking.puml",
    25: "create-walk-in-booking.puml",
    26: "view-booking-detail.puml",
    27: "view-my-bookings.puml",
    28: "view-booking-calendar.puml",
    29: "confirm-booking.puml",
    30: "reject-booking.puml",
    31: "reschedule-booking.puml",
    32: "preview-cancellation-fee-refund.puml",
    33: "cancel-booking.puml",
    34: "check-in-booking.puml",
    35: "complete-booking.puml",
    36: "mark-no-show.puml",
    37: "handle-booking-conflict.puml",
    38: "view-checkout-summary.puml",
    39: "choose-payment-option.puml",
    40: "40.pay-online-sandbox.puml",
    41: "41.capture-online-payment.puml",
    42: "confirm-remaining-payment.puml",
    43: "handle-failed-expired-payment.puml",
    44: "view-payment-history.puml",
    45: "generate-booking-invoice.puml",
    46: "view-invoice-payment-status.puml",
    47: "process-online-refund.puml",
    48: "manage-refund-requests.puml",
    49: "configure-deposit-rules.puml",
    50: "configure-cancellation-refund-policy.puml",
    51: "view-active-promotions.puml",
    52: "manage-promotion-campaigns.puml",
    53: "apply-promotion-to-booking.puml",
    54: "view-membership-progress.puml",
    55: "manage-membership-rules.puml",
    56: "view-membership-benefits.puml",
    57: "send-booking-confirmation-notification.puml",
    58: "send-booking-reminder-notification.puml",
    59: "send-cancellation-refund-notification.puml",
    60: "view-revenue-report.puml",
    61: "view-booking-report.puml",
    62: "view-customer-activity-report.puml",
    63: "ask-smart-assistant-for-available-fields.puml",
    64: "get-suggested-available-slots.puml",
}


SEQUENCE_FILES = {
    1: "Register-and-Verify-Email.puml",
    2: "Login.puml",
    **{n: next(p.name for p in SEQUENCE_DIR.glob(f"UC-{n:02d}-*.puml")) for n in range(3, 65)},
}


def paragraph_has_drawing(paragraph: Paragraph) -> bool:
    return bool(paragraph._p.xpath(".//w:drawing|.//w:pict"))


def remove_paragraph(paragraph: Paragraph) -> None:
    element = paragraph._element
    element.getparent().remove(element)


def insert_paragraph_after(paragraph: Paragraph) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    return Paragraph(new_p, paragraph._parent)


def fitted_size(image_path: Path, max_width: float = 6.35, max_height: float = 8.0):
    with Image.open(image_path) as image:
        width_px, height_px = image.size
    scale = min(max_width / width_px, max_height / height_px)
    return Inches(width_px * scale), Inches(height_px * scale)


def add_diagram_after(heading: Paragraph, image_path: Path, alt_text: str, *, max_height: float = 8.0) -> None:
    paragraph = insert_paragraph_after(heading)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Inches(0.04)
    paragraph.paragraph_format.space_after = Inches(0.08)
    paragraph.paragraph_format.keep_together = True
    width, height = fitted_size(image_path, max_height=max_height)
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=width, height=height)
    for doc_pr in paragraph._p.xpath(".//wp:docPr"):
        doc_pr.set("name", alt_text)
        doc_pr.set("title", alt_text)
        doc_pr.set("descr", alt_text)


def heading_map(document: Document) -> dict[int, dict[str, Paragraph]]:
    result: dict[int, dict[str, Paragraph]] = {}
    current_uc: int | None = None
    for paragraph in document.paragraphs:
        text = " ".join(paragraph.text.split())
        if paragraph.style.name == "Heading 2":
            match = re.match(r"(\d+)\.", text)
            current_uc = int(match.group(1)) if match else None
            if current_uc is not None:
                result.setdefault(current_uc, {})["use_case"] = paragraph
        elif current_uc is not None and paragraph.style.name == "Heading 3":
            lowered = text.lower()
            if lowered.startswith("a. class diagram"):
                result[current_uc]["class"] = paragraph
            elif lowered.startswith("c. sequence diagram"):
                result[current_uc]["sequence"] = paragraph
    return result


def remove_existing_diagrams(document: Document, headings: dict[int, dict[str, Paragraph]]) -> int:
    diagram_headings = {
        item[k]._p
        for uc, item in headings.items()
        if uc <= 64
        for k in ("class", "sequence")
        if k in item
    }
    active = False
    removed = 0
    for paragraph in list(document.paragraphs):
        if paragraph._p in diagram_headings:
            active = True
            continue
        if paragraph.style.name.startswith("Heading"):
            active = False
        if active and paragraph_has_drawing(paragraph):
            remove_paragraph(paragraph)
            removed += 1
    return removed


def diagram_exists_before_next_heading(document: Document, heading: Paragraph) -> bool:
    paragraphs = document.paragraphs
    start = next(i for i, paragraph in enumerate(paragraphs) if paragraph._p is heading._p)
    for paragraph in paragraphs[start + 1 :]:
        if paragraph.style.name.startswith("Heading"):
            return False
        if paragraph_has_drawing(paragraph):
            return True
    return False


def set_update_fields(document: Document) -> None:
    settings = document.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("input_docx", type=Path)
    parser.add_argument("output_docx", type=Path)
    parser.add_argument("--class-png-dir", type=Path, required=True)
    parser.add_argument("--sequence-png-dir", type=Path, required=True)
    parser.add_argument("--package-png", type=Path)
    parser.add_argument(
        "--preserve-existing",
        action="store_true",
        help="Keep existing UC diagram images and insert only missing diagrams.",
    )
    args = parser.parse_args()

    document = Document(args.input_docx)
    if args.package_png:
        if not args.package_png.is_file():
            raise FileNotFoundError(args.package_png)
        package_heading = next(
            paragraph for paragraph in document.paragraphs
            if " ".join(paragraph.text.split()) == "1. Code Packages"
        )
        following = insert_paragraph_after(package_heading)
        following.alignment = WD_ALIGN_PARAGRAPH.CENTER
        following.paragraph_format.space_before = Inches(0.04)
        following.paragraph_format.space_after = Inches(0.08)
        following.paragraph_format.keep_together = True
        width, height = fitted_size(args.package_png, max_height=4.8)
        run = following.add_run()
        run.add_picture(str(args.package_png), width=width, height=height)
        for doc_pr in following._p.xpath(".//wp:docPr"):
            doc_pr.set("name", "Code Package Diagram")
            doc_pr.set("title", "Code Package Diagram")
            doc_pr.set("descr", "Code-first Java package dependencies")
        for paragraph in list(document.paragraphs):
            if " ".join(paragraph.text.split()) == "Package diagram will be inserted here during the diagram phase.":
                remove_paragraph(paragraph)
                break
    headings = heading_map(document)
    missing = [n for n in range(1, 65) if not {"class", "sequence"}.issubset(headings.get(n, {}))]
    if missing:
        raise RuntimeError(f"SDS diagram headings missing for UCs: {missing}")

    removed = 0 if args.preserve_existing else remove_existing_diagrams(document, headings)
    inserted = 0
    preserved = 0
    for uc in range(1, 65):
        class_png = args.class_png_dir / Path(CLASS_FILES[uc]).with_suffix(".png").name
        sequence_png = args.sequence_png_dir / Path(SEQUENCE_FILES[uc]).with_suffix(".png").name
        for path in (class_png, sequence_png):
            if not path.is_file():
                raise FileNotFoundError(path)
        for kind, image_path, alt_text in (
            ("class", class_png, f"UC-{uc:02d} Class Diagram"),
            ("sequence", sequence_png, f"UC-{uc:02d} Sequence Diagram"),
        ):
            heading = headings[uc][kind]
            if args.preserve_existing and diagram_exists_before_next_heading(document, heading):
                preserved += 1
                continue
            add_diagram_after(heading, image_path, alt_text)
            inserted += 1

    set_update_fields(document)
    args.output_docx.parent.mkdir(parents=True, exist_ok=True)
    document.save(args.output_docx)
    print(f"Removed {removed} existing UC diagram images")
    print(f"Preserved {preserved} existing UC diagram images")
    print(f"Inserted {inserted} code-first diagram images")
    print(args.output_docx)


if __name__ == "__main__":
    main()
