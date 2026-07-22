#!/usr/bin/env python3
"""Replace the database-schema image in a template-derived DOCX."""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Inches
from docx.text.paragraph import Paragraph
from PIL import Image


def insert_paragraph_after(paragraph: Paragraph) -> Paragraph:
    node = OxmlElement("w:p")
    paragraph._p.addnext(node)
    return Paragraph(node, paragraph._parent)


def has_drawing(paragraph: Paragraph) -> bool:
    return bool(paragraph._p.xpath(".//w:drawing|.//w:pict"))


def fitted_size(path: Path, max_width: float = 6.35, max_height: float = 8.0):
    with Image.open(path) as image:
        width, height = image.size
    scale = min(max_width / width, max_height / height)
    return Inches(width * scale), Inches(height * scale)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("input_docx", type=Path)
    parser.add_argument("output_docx", type=Path)
    parser.add_argument("erd_png", type=Path)
    args = parser.parse_args()

    document = Document(args.input_docx)
    heading = next(
        paragraph for paragraph in document.paragraphs
        if " ".join(paragraph.text.split()).lower() == "a. database schema"
    )
    paragraphs = document.paragraphs
    start = next(index for index, paragraph in enumerate(paragraphs) if paragraph._p is heading._p)
    removed = 0
    for paragraph in paragraphs[start + 1:]:
        if paragraph.style.name.startswith("Heading"):
            break
        if has_drawing(paragraph):
            node = paragraph._element
            node.getparent().remove(node)
            removed += 1

    diagram = insert_paragraph_after(heading)
    diagram.alignment = WD_ALIGN_PARAGRAPH.CENTER
    diagram.paragraph_format.space_before = Inches(0.04)
    diagram.paragraph_format.space_after = Inches(0.08)
    diagram.paragraph_format.keep_together = True
    width, height = fitted_size(args.erd_png)
    run = diagram.add_run()
    run.add_picture(str(args.erd_png), width=width, height=height)
    for properties in diagram._p.xpath(".//wp:docPr"):
        properties.set("name", "GoalZone Code-First Database Schema")
        properties.set("title", "GoalZone Code-First Database Schema")
        properties.set("descr", "PostgreSQL schema aligned with Flyway V1 through V11")

    args.output_docx.parent.mkdir(parents=True, exist_ok=True)
    document.save(args.output_docx)
    print(f"Replaced {removed} database-schema image(s): {args.output_docx}")


if __name__ == "__main__":
    main()
