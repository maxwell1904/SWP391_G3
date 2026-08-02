from __future__ import annotations

from copy import deepcopy
from pathlib import Path
import re

from PIL import Image
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.table import Table
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[2]
PROJECT = ROOT / "SWP391_Project"
IMPORTS = ROOT / "docs_new_imports"
OUTPUT = PROJECT / "output" / "doc"
CLASS_RENDER = ROOT / "tmp" / "class_render_v17"
SEQUENCE_RENDER = ROOT / "tmp" / "sequence_render_v17"
PACKAGE_IMAGE = PROJECT / "docs" / "generated" / "package" / "code-packages.png"

RDS_INPUT = IMPORTS / "RDS Document_G03_Final.docx"
SDS_INPUT = IMPORTS / "SDS Document_G03_Final.docx"
RDS_OUTPUT = OUTPUT / "GoalZone_RDS_Retake_Aligned_2026-08-01.docx"
SDS_OUTPUT = OUTPUT / "GoalZone_SDS_Retake_Aligned_2026-08-01.docx"

FINAL_MODULES = (
    ("Account and Access", 1, 10),
    ("Fields, Slots, Services, and Issues", 11, 23),
    ("Booking and Venue Operations", 24, 33),
    ("Checkout, Payments, Refunds, and Policy", 34, 44),
    ("Promotions, Membership, Notifications, and Reports", 45, 56),
)

# (source old ID, feature, owner, actor, final name, catalogue description)
FINAL_USE_CASES = (
    (1, "Account onboarding", "BonVT", "Guest", "Register account", "A Guest creates a Customer account and completes email verification."),
    (2, "Authentication", "BonVT", "Guest/User", "Login", "An active user signs in with email or phone and receives a revocable JWT."),
    (3, "Authentication", "BonVT", "User", "Logout", "A signed-in user signs out and invalidates all previously issued JWTs for that account."),
    (4, "Profile management", "BonVT", "User", "Manage personal profile", "A user views and updates only their own editable profile information."),
    (5, "Credential security", "BonVT", "User", "Change password", "A signed-in user replaces their password after proving the current password."),
    (6, "Account recovery", "BonVT", "Guest/User", "Recover password", "A Guest requests, validates, and consumes a one-hour single-use password recovery link."),
    (7, "Customer account administration", "BonVT", "Admin", "View customer accounts", "An Admin lists Customer accounts and opens business-facing account and activity context in read-only mode."),
    (7, "Customer account administration", "BonVT", "Admin", "Lock/unlock customer account", "An Admin locks or restores Customer sign-in access with a reason, token revocation, and notification."),
    (8, "Staff account administration", "BonVT", "Admin", "Manage staff accounts", "An Admin creates and maintains Staff identity/status while Staff privately owns the invited password."),
    (9, "Customer activity", "BonVT", "Admin", "View customer activity", "An Admin opens an operational activity drill-down for one Customer using business context rather than raw IDs."),
    (10, "Field catalogue", "BaoNG", "Guest/Customer", "View field list", "A visitor views active football fields and their catalogue context."),
    (11, "Field catalogue", "BaoNG", "Guest/Customer", "View field detail", "A visitor views the selected field, images, facilities, prices, and active services."),
    (12, "Availability search", "BaoNG", "Guest/Customer", "Search available fields by date/time", "A visitor searches generated, priced, non-blocked field slots whose start time is still in the future, with optional filters."),
    (13, "Field management", "BaoNG", "Admin", "Manage football fields", "An Admin creates, updates, activates, deactivates, and uploads images for football fields."),
    (14, "Field pricing", "BaoNG", "Admin", "Manage field pricing by time range", "An Admin maintains effective weekday/weekend price bands without any fallback price."),
    (15, "Slot generation rules", "BaoNG", "Admin", "Configure slot generation rules", "An Admin configures opening time, closing time, duration, and the rolling generation horizon."),
    (15, "Slot exceptions", "BaoNG", "Staff", "Block/unblock field slots", "Venue Staff records or removes operational exceptions without overwriting active bookings."),
    (16, "Operations calendar", "BaoNG", "Staff", "View field operation calendar", "Venue Staff views generated available, blocked, and booked field slots for one date."),
    (17, "Extra-service catalogue", "BaoNG", "Admin", "Manage extra services", "An Admin maintains priced rental, sale, and staff services and their limits/status."),
    (19, "Booking add-ons", "BaoNG", "Customer/Staff", "Add extra services to booking", "Customer or Staff selects validated extra services while creating a booking."),
    (20, "Booking add-ons", "BaoNG", "Customer/Staff", "Update extra services before check-in", "The booking owner or Venue Staff edits add-ons before start and the system recalculates all financial effects."),
    (21, "Issue management", "BaoNG", "Customer/Staff", "Report field/service issue", "An authenticated Customer or Staff member reports a described issue under their own identity."),
    (22, "Issue management", "BaoNG", "Staff", "Resolve field/service issue", "Venue Staff progresses an open issue to a terminal resolved or rejected state with notes."),
    (23, "Booking creation", "NgocPA", "Customer", "Create online booking", "A Customer creates a validated pending online booking and continues to PayPal Sandbox."),
    (24, "Booking creation", "NgocPA", "Staff", "Create walk-in booking", "Venue Staff creates a walk-in booking for either a matched Customer account or a first-time visitor identified by name and phone, then records cash or leaves payment pending."),
    (25, "Booking details", "NgocPA", "Customer/Staff", "View booking detail", "An authorized user views slot, service, promotion, payment, invoice, and refund detail for a booking."),
    (26, "Booking history", "NgocPA", "Customer", "View my bookings", "A Customer views current and past owned bookings."),
    (27, "Booking calendar", "NgocPA", "Staff", "View booking calendar", "Venue Staff views the daily booking queue by field and date."),
    (30, "Booking rescheduling", "NgocPA", "Customer/Staff", "Reschedule booking", "Customer or Staff moves an eligible booking before start and the system recalculates the destination and balance."),
    (32, "Booking cancellation", "NgocPA", "Customer/Staff", "Cancel booking", "Customer or Staff reviews policy amounts and cancels an eligible booking before it starts."),
    (33, "Venue operations", "NgocPA", "Staff", "Check in booking", "Venue Staff checks in a paid confirmed booking from 30 minutes before its start."),
    (34, "Venue operations", "NgocPA", "Staff", "Complete booking", "Venue Staff completes a started checked-in booking only after the remaining balance is zero."),
    (35, "Venue operations", "NgocPA", "Staff", "Mark no-show", "Venue Staff records no-show from 15 minutes after start without increasing membership progress."),
    (37, "Checkout", "AnNP", "Customer/Staff", "View checkout summary", "Customer or Staff reviews authoritative field, service, discount, deposit, paid, and remaining amounts."),
    (38, "Checkout", "AnNP", "Customer/Staff", "Choose payment option", "A Customer chooses online deposit/full payment; Staff chooses an eligible counter-cash amount."),
    (39, "Online payment", "AnNP", "Customer", "Pay online via PayPal Sandbox", "A Customer approves a PayPal Sandbox order whose verified idempotent capture automatically confirms eligible booking state."),
    (41, "Counter payment", "AnNP", "Staff", "Confirm remaining payment", "Venue Staff records a real on-site cash payment against a booking balance."),
    (42, "Payment recovery", "AnNP", "Customer/PayPal/Scheduler", "Handle failed/expired payment", "Provider failure remains visible/retryable and the scheduler releases aged unpaid holds."),
    (43, "Payment records", "AnNP", "Customer/Staff/Admin", "View payment history", "An authorized user views persisted cash/PayPal transactions with booking, field, slot, amount, time, and refund context."),
    (45, "Billing status", "AnNP", "Customer/Staff", "View invoice/payment status", "An authorized user views the system-maintained invoice, paid, remaining, refunded, and transaction status."),
    (46, "Refund management", "AnNP", "Staff/PayPal", "Process online refund", "Venue Staff submits an approved online refund idempotently and accepts only provider COMPLETED as final."),
    (47, "Refund management", "AnNP", "Customer/Staff", "Manage refund requests", "The booking owner submits an amount and reason; Venue Staff reviews its auditable lifecycle while Admin remains read-only."),
    (48, "Booking policies", "AnNP", "Admin", "Configure deposit rules", "An Admin configures the default online deposit percentage."),
    (49, "Booking policies", "AnNP", "Admin", "Configure cancellation/refund policy", "An Admin configures before-24-hour and same-day refund percentages."),
    (50, "Promotions", "AnPTT", "Customer/Guest", "View active promotions", "A visitor views currently eligible active promotion campaigns."),
    (51, "Promotions", "AnPTT", "Admin", "Manage promotion campaigns", "An Admin creates, updates, activates, deactivates, and uploads banners for campaigns."),
    (52, "Promotions", "AnPTT", "Customer/System", "Apply promotion to booking", "A Customer applies one eligible code whose discount and reserved usage are snapshotted."),
    (53, "Membership", "AnPTT", "Customer", "View membership progress", "A Customer views progress under the active lifetime, monthly, or consecutive-week rule."),
    (54, "Membership", "AnPTT", "Admin", "Manage membership rules", "An Admin maintains qualification periods, thresholds, ordering, status, discounts, and benefits."),
    (55, "Membership", "AnPTT", "Customer/Guest", "View membership benefits", "A visitor views active membership levels and benefit descriptions."),
    (56, "Notifications", "AnPTT", "Customer", "Send booking confirmation notification", "A confirmation event creates one visible notification for the affected Customer."),
    (57, "Notifications", "AnPTT", "Scheduler/Customer", "Send booking reminder notification", "The scheduler creates at most one reminder inside the configured pre-start window."),
    (58, "Notifications", "AnPTT", "Customer", "Send cancellation/refund notification", "Cancellation and completed refund events notify the affected Customer."),
    (59, "Reporting", "AnPTT", "Admin", "View revenue report", "An Admin views gross collected, fees, refunds, net revenue, discounts, and commercial values for a period."),
    (60, "Reporting", "AnPTT", "Admin", "View booking report", "An Admin views valid lifecycle counts, peak time, and field utilization for a period."),
    (61, "Reporting", "AnPTT", "Admin", "View customer activity report", "An Admin views aggregate returning/top-customer and membership activity for a selected period."),
)

assert len(FINAL_USE_CASES) == 56

# Alternative flows are meaningful business branches, not a quota of extra steps.
# A one-step branch is retained when one step completely describes the variation.
ALTERNATIVE_FLOW_OVERRIDES = {
    1: "A1: Email delivery is delayed\n1. The account remains unverified and the Customer can request a new verification email.",
    2: "A1: Sign in by phone\n1. The user supplies the registered phone number instead of email and continues through the same credential check.",
    3: "A1: Server logout cannot be reached\n1. The client still clears its local token; server-side tokens remain subject to expiry until a later successful revocation.",
    4: "A1: No profile changes\n1. The user leaves without saving and the stored profile remains unchanged.",
    5: "A1: New password matches the current password\n1. GoalZone asks the user to choose a different password.",
    6: "A1: Earlier reset link exists\n1. Issuing a new link invalidates the earlier unused link.",
    7: "A1: No activity\n1. The Customer detail opens with a clear empty state instead of raw identifiers.",
    8: "A1: Unlock account\n1. Admin confirms restoration; GoalZone activates the account, clears the reason, revokes older tokens, and notifies the Customer.",
    9: "A1: Edit existing Staff account\n1. Admin updates identity or status without viewing or replacing the Staff member's password.",
    10: "A1: No activity in the selected context\n1. GoalZone shows an empty operational summary and leaves the Customer account unchanged.",
    11: "A1: No active fields\n1. The catalogue displays an explicit no-fields message.",
    12: "A1: Optional services are unavailable\n1. The field detail remains viewable and states that no active add-ons are configured.",
    13: "A1: No date supplied\n1. GoalZone uses tomorrow as the default search date.\nA2: No available result\n1. Booked and blocked rows remain unavailable and the UI shows that no selectable slots match.",
    14: "A1: Deactivate field\n1. Admin confirms deactivation; the field remains in administration history but disappears from public selection.",
    15: "A1: Retire a price band\n1. Admin marks the rule inactive; other effective non-overlapping rules continue to price slots.",
    16: "A1: Scheduled horizon extension\n1. The scheduler applies the saved rules to extend the rolling calendar without changing those rules.",
    17: "A1: Future exception outside generated rows\n1. Staff records the exception and later slot generation preserves it.\nA2: Restore blocked slot\n1. Staff selects Unblock and GoalZone makes the row available again when no booking occupies it.",
    18: "A1: Empty operating day\n1. The calendar shows generated available rows even when no bookings or blocks exist.",
    19: "A1: Deactivate an existing service\n1. Admin retires the service from new selection while preserving it on historical bookings.",
    20: "A1: No add-on selected\n1. Booking continues with field price only.",
    21: "A1: Remove all add-ons\n1. The user saves zero quantities and GoalZone recalculates the invoice without service charges.",
    22: "A1: General issue\n1. The reporter leaves booking and field blank and submits a venue-level issue.",
    23: "A1: Reject report\n1. Staff records a reason and closes the report as rejected rather than resolved.",
    24: "A1: Pay full amount\n1. The Customer chooses full payment instead of the default deposit before PayPal approval.",
    25: "A1: First-time visitor has no account\n1. Staff selects First-time visitor after no account matches the supplied phone or email.\n2. Staff enters the visitor's name, phone, and optional email.\n3. GoalZone stores those details on the booking without creating an account, then continues at Normal Flow step 3.\nA2: Pay later\n1. Staff chooses Pay later instead of immediate cash.\n2. GoalZone creates a pending booking with the full balance remaining and no payment row.",
    26: "A1: No payments or refunds yet\n1. The booking detail shows the invoice with explicit empty transaction sections.",
    27: "A1: No bookings\n1. The Customer sees an empty state with a direct action to start a booking.",
    28: "A1: Filter to past records\n1. Staff selects a previous operating date and reviews historical bookings without changing them.\nA2: No bookings on the date\n1. Staff sees the field schedule and an empty booking queue for that day.",
    29: "A1: Destination costs more\n1. GoalZone preserves paid value and shows the new remaining balance.\nA2: Destination costs less\n1. GoalZone creates refundable value for the difference without auto-paying it.",
    30: "A1: Cancellation has no refund\n1. After preview, the user may still confirm cancellation and GoalZone records a zero refundable amount.",
    31: "A1: Staff adds an audit note\n1. The note is stored with the confirmed check-in action.",
    32: "A1: Remaining balance exists\n1. Staff records the remaining payment first, then confirms completion.",
    33: "A1: Customer arrives before confirmation\n1. Staff cancels the no-show action and uses Check-in instead.\nA2: No-show confirmed\n1. Staff must enter an operational note before GoalZone records no-show.",
    34: "A1: Promotion is not eligible\n1. GoalZone keeps the base checkout amounts and explains why the code was not applied.",
    35: "A1: Staff chooses pay later\n1. Checkout displays the full unpaid balance and preserves pending status.\nA2: Customer chooses full payment\n1. Checkout sends the total amount instead of the deposit amount.",
    36: "A1: Customer cancels PayPal approval\n1. GoalZone releases the unpaid hold and records no captured payment.",
    37: "A1: Partial cash is not accepted\n1. Staff chooses the supported remaining-balance action or leaves the balance unchanged.",
    38: "A1: Customer retries before expiry\n1. A new provider attempt is associated with the same eligible pending booking.\nA2: Hold expires\n1. The scheduler expires the unpaid booking and releases the slot.",
    39: "A1: No payment records\n1. The history displays an empty state rather than a booking code without context.",
    40: "A1: Refund is still processing\n1. The invoice keeps paid, remaining, refundable, and refund status visible without treating the refund as completed.",
    41: "A1: Provider reports pending\n1. The refund remains processing and Staff may retry safely with the same idempotency key.",
    42: "A1: Staff rejects request\n1. Staff supplies a review note and GoalZone closes the request as rejected.\nA2: Admin opens the case\n1. Admin may inspect the audit trail but receives no processing action.",
    43: "A1: Deposit percentage is changed\n1. New checkout previews use the new percentage; stored booking snapshots remain unchanged.",
    44: "A1: Same-day cancellation\n1. GoalZone applies the configured same-day percentage rather than the early-cancellation percentage.",
    45: "A1: Visitor chooses an offer\n1. GoalZone copies the selected campaign code into booking checkout for eligibility validation.",
    46: "A1: Campaign is not stackable\n1. Admin leaves stacking off so an eligible promotion replaces the membership discount.\nA2: Deactivate campaign\n1. Admin confirms deactivation and the campaign stops appearing to visitors.",
    47: "A1: Code is eligible but capacity is exhausted\n1. GoalZone rejects the code and leaves checkout undiscounted.\nA2: Booking is cancelled before use\n1. Reserved campaign usage is released.",
    48: "A1: Monthly or weekly rule is active\n1. Progress uses the active period metric while lifetime completed bookings remain informational only.",
    49: "A1: Deactivate a tier\n1. Admin confirms deactivation; the tier remains historical but is excluded from new qualification.",
    50: "A1: Visitor is not signed in\n1. Active tier benefits remain visible without showing personal progress.",
    51: "A1: Notification already exists\n1. GoalZone keeps the single confirmation notification and does not create a duplicate.",
    52: "A1: Reminder was already sent\n1. The scheduler skips the booking to preserve at-most-once reminder delivery.",
    53: "A1: Refund fails or remains processing\n1. GoalZone reports the current refund state and sends completion wording only after final success.",
    54: "A1: No date range\n1. Admin views all-time revenue; payments use paid time and refunds use processed time.",
    55: "A1: No date range\n1. Admin views all-time booking metrics based on booked slot date.",
    56: "A1: No qualifying Customer activity\n1. The report returns zero totals and empty ranked lists without exposing raw identifiers.",
}

assert len(ALTERNATIVE_FLOW_OVERRIDES) == 56

CLASS_IMAGE_NAMES = (
    "register.png", "login.png", "logout.png", "manage-personal-profile.png", "change-password.png",
    "forgot-reset-password.png", "manage-customer-accounts.png", "lock-unlock-customer-account.png",
    "manage-staff-accounts.png", "view-customer-activity.png", "view-field-list.png", "view-field-detail.png",
    "search-available-fields.png", "manage-football-fields.png", "manage-field-pricing-by-time-range.png",
    "configure-slot-generation-rules.png", "block-unblock-field-slots.png", "view-field-operation-calendar.png",
    "manage-extra-services.png", "add-extra-services-to-booking.png", "update-extra-services-before-check-in.png",
    "report-field-service-issue.png", "resolve-field-service-issue.png", "create-online-booking.png",
    "create-walk-in-booking.png", "view-booking-detail.png", "view-my-bookings.png", "view-booking-calendar.png",
    "reschedule-booking.png", "cancel-booking.png", "check-in-booking.png", "complete-booking.png",
    "mark-no-show.png", "view-checkout-summary.png", "choose-payment-option.png", "pay-online-sandbox.png",
    "confirm-remaining-payment.png", "handle-failed-expired-payment.png", "view-payment-history.png",
    "view-invoice-payment-status.png", "process-online-refund.png", "manage-refund-requests.png",
    "configure-deposit-rules.png", "configure-cancellation-refund-policy.png", "view-active-promotions.png",
    "manage-promotion-campaigns.png", "apply-promotion-to-booking.png", "view-membership-progress.png",
    "manage-membership-rules.png", "view-membership-benefits.png", "send-booking-confirmation-notification.png",
    "send-booking-reminder-notification.png", "send-cancellation-refund-notification.png", "view-revenue-report.png",
    "view-booking-report.png", "view-customer-activity-report.png",
)

assert len(CLASS_IMAGE_NAMES) == 56

OLD_REFERENCE_TARGETS = {
    7: "UC-07/UC-08", 8: "UC-09", 9: "UC-10", 10: "UC-11", 11: "UC-12", 12: "UC-13",
    13: "UC-14", 14: "UC-15", 15: "UC-16/UC-17", 16: "UC-18", 17: "UC-19", 18: "UC-20",
    19: "UC-20", 20: "UC-21", 21: "UC-22", 22: "UC-23", 23: "UC-24", 24: "UC-25",
    25: "UC-26", 26: "UC-27", 27: "UC-28", 28: "UC-36", 29: "UC-24", 30: "UC-29",
    31: "UC-30", 32: "UC-30", 33: "UC-31", 34: "UC-32", 35: "UC-33", 36: "UC-29/UC-30",
    37: "UC-34", 38: "UC-35", 39: "UC-36", 40: "UC-36", 41: "UC-37", 42: "UC-38",
    43: "UC-39", 44: "UC-40", 45: "UC-40", 46: "UC-41", 47: "UC-42", 48: "UC-43",
    49: "UC-44", 50: "UC-45", 51: "UC-46", 52: "UC-47", 53: "UC-48", 54: "UC-49",
    55: "UC-50", 56: "UC-51", 57: "UC-52", 58: "UC-53", 59: "UC-54", 60: "UC-55",
    61: "UC-56", 62: "UC-13", 63: "UC-13", 64: "UC-13",
}


def iter_blocks(document: Document):
    for child in document.element.body.iterchildren():
        if child.tag.endswith("}p"):
            yield Paragraph(child, document)
        elif child.tag.endswith("}tbl"):
            yield Table(child, document)


def iter_all_paragraphs(document: Document):
    yield from document.paragraphs
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs
                for nested in cell.tables:
                    for nested_row in nested.rows:
                        for nested_cell in nested_row.cells:
                            yield from nested_cell.paragraphs


def replace_across_runs(paragraph: Paragraph, old: str, new: str) -> int:
    count = 0
    while old in paragraph.text:
        full = "".join(run.text for run in paragraph.runs)
        start = full.find(old)
        if start < 0:
            break
        end = start + len(old)
        positions = []
        cursor = 0
        for index, run in enumerate(paragraph.runs):
            positions.append((index, cursor, cursor + len(run.text)))
            cursor += len(run.text)
        start_info = next((x for x in positions if x[1] <= start < x[2]), None)
        end_info = next((x for x in positions if x[1] < end <= x[2]), None)
        if start_info is None or end_info is None:
            break
        start_run, start_lo, _ = start_info
        end_run, end_lo, _ = end_info
        start_offset = start - start_lo
        end_offset = end - end_lo
        if start_run == end_run:
            text = paragraph.runs[start_run].text
            paragraph.runs[start_run].text = text[:start_offset] + new + text[end_offset:]
        else:
            first_text = paragraph.runs[start_run].text
            last_text = paragraph.runs[end_run].text
            paragraph.runs[start_run].text = first_text[:start_offset] + new
            for index in range(start_run + 1, end_run):
                paragraph.runs[index].text = ""
            paragraph.runs[end_run].text = last_text[end_offset:]
        count += 1
    return count


def replace_text(document: Document, replacements: list[tuple[str, str]]) -> None:
    for paragraph in iter_all_paragraphs(document):
        for old, new in replacements:
            replace_across_runs(paragraph, old, new)


def update_cover_metadata(document: Document) -> None:
    """Normalize final-submission metadata while preserving template styling."""
    replace_text(
        document,
        [
            ("Can Tho, May 2026", "Can Tho, August 2026"),
            ("Record of changeS", "Record of Changes"),
        ],
    )
    for paragraph in document.paragraphs:
        if paragraph.text.strip() == "af":
            for run in paragraph.runs:
                run.text = ""


def replace_xml_attribute_text(document: Document, replacements: list[tuple[str, str]]) -> None:
    """Keep non-visible image descriptions and other OOXML metadata aligned."""
    for element in document.element.iter():
        for attribute, value in list(element.attrib.items()):
            updated = value
            for old, new in replacements:
                updated = updated.replace(old, new)
            if updated != value:
                element.set(attribute, updated)


def has_drawing(paragraph: Paragraph) -> bool:
    return bool(paragraph._p.xpath(".//w:drawing | .//w:pict"))


def image_fit(path: Path, max_width: float = 6.35, max_height: float = 7.0):
    with Image.open(path) as image:
        width, height = image.size
    scale = min(max_width / width, max_height / height)
    return Inches(width * scale), Inches(height * scale)


def replace_image(paragraph: Paragraph, path: Path, max_width: float = 6.35, max_height: float = 7.0) -> None:
    for run in list(paragraph.runs):
        paragraph._p.remove(run._r)
    width, height = image_fit(path, max_width, max_height)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    # Some imported template picture paragraphs use compressed spacing. Reset it
    # so a tall replacement diagram cannot overlap the preceding subsection label.
    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.add_run().add_picture(str(path), width=width, height=height)


def replace_drawing_after(
    document: Document,
    section_heading: str,
    marker: str,
    image_path: Path,
    max_height: float = 7.0,
) -> None:
    blocks = list(iter_blocks(document))
    in_section = False
    for index, block in enumerate(blocks):
        if isinstance(block, Paragraph) and block.style.name.startswith("Heading 2"):
            in_section = block.text.strip() == section_heading
        if not in_section or not isinstance(block, Paragraph) or block.text.strip() != marker:
            continue
        for candidate in blocks[index + 1 : index + 5]:
            if isinstance(candidate, Paragraph) and has_drawing(candidate):
                replace_image(candidate, image_path, max_height=max_height)
                return
        raise RuntimeError(f"Image not found after {section_heading!r} / {marker!r}")
    raise RuntimeError(f"Section not found: {section_heading!r}")


def remove_rds_external_diagram_section(
    document: Document,
    start_heading: str,
    end_heading: str,
) -> None:
    """Remove a provisional diagram section until the approved external template is supplied."""
    body = document.element.body
    blocks = list(body.iterchildren())
    start = next(
        i for i, element in enumerate(blocks)
        if element.tag.endswith("}p")
        and Paragraph(element, document).text.strip() == start_heading
    )
    end = next(
        i for i, element in enumerate(blocks)
        if element.tag.endswith("}p")
        and Paragraph(element, document).text.strip() == end_heading
    )
    for element in blocks[start:end]:
        body.remove(element)

    target = Paragraph(blocks[end], document)
    for drawing in list(target._p.xpath(".//w:drawing | .//w:pict")):
        drawing.getparent().remove(drawing)
    target.paragraph_format.page_break_before = False


def align_rds_retake_rules(document: Document) -> None:
    replacements = [
        ("a. Usecase Diagrams", "a. Use Case Diagrams"),
        ("Field pricing uses the first active matching day/time rule or USD 12 fallback.",
         "A slot is bookable only when an active day/time pricing rule matches; the system never invents a fallback price."),
        ("A2: No matching rule\n1. Checkout calculation uses the hard-coded USD 12 fallback.",
         "A2: No matching rule\n1. The slot is excluded from bookable results and checkout is rejected until Admin configures a valid price."),
        ("If no active price band matches, the code-first fallback is USD 12.",
         "Every bookable slot must match an active pricing rule."),
        ("It uses findFirst without explicit ordering and does not apply effective dates.",
         "It applies effective-from/effective-to dates; overlapping active rules are rejected when configured."),
        ("PRE-2: A Customer owns the booking, or the caller is Staff/Admin.",
         "PRE-2: A Customer owns the booking, or the caller is Venue Staff."),
        ("POST-2: Staff/Admin can see it in the issue list.",
         "POST-2: Venue Staff can handle it; Admin may inspect it as audit context."),
        ("PRE-2: The issue title is provided.", "PRE-2: The issue title and description are provided."),
        ("A1: General issue\n1. No booking, field or service association is supplied; the issue is still accepted.\nA2: Operator report/assignment\n1. Staff/Admin may report against any booking and may assign the issue to an existing Staff user.",
         "A1: General issue\n1. A Customer or Venue Staff member may report a general issue without a booking link.\nA2: Venue report\n1. Venue Staff may report against a booking and is assigned from the authenticated identity."),
        ("EXC-3: Only a Staff-role user can be assigned; an Admin account cannot be the assignee.",
         "EXC-3: Closed issues cannot be reopened, and the browser cannot choose an assignee identity."),
        ("EXC-3: An assignment target whose role is not Staff is rejected.\nEXC-4: The backend permits blank description and no association, although the React Customer/Staff forms require a description.",
         "EXC-3: Blank title or description is rejected.\nEXC-4: Booking, field, and service links must agree with the selected booking."),
        ("Implementation rule: Reporter identity is always the authenticated principal. Only operators may assign at creation. Customers cannot retrieve the global issue list.",
         "Implementation rule: Reporter identity is always the authenticated principal. Venue Staff is self-assigned; Customers cannot retrieve the global issue list."),
        ("PRE-2: The caller has Staff or Admin authority.", "PRE-2: The caller is Venue Staff."),
        ("A1: Reject issue\n1. An API operator submits rejected with a required resolution note.\n2. The issue is closed and the reporter is notified.\nA2: Reopen/change state\n1. An operator submits open or in_progress.\n2. resolvedAt is cleared; there is no transition guard.",
         "A1: Reject issue\n1. Venue Staff submits rejected with a required resolution note.\n2. The issue is closed and the reporter is notified.\nA2: Start work\n1. Venue Staff moves an open issue to in_progress.\n2. Only resolved or rejected may follow; closed states are terminal."),
        ("Implementation rule: Statuses are open, in_progress, resolved and rejected. The service permits transitions between any statuses; the React Staff UI exposes Start and Resolve only.",
         "Implementation rule: Statuses are open, in_progress, resolved and rejected. The service enforces open -> in_progress/resolved/rejected and in_progress -> resolved/rejected."),
        ("EXC-1: Staff/Admin cannot create online bookings; Customers cannot create walk-ins.",
         "EXC-1: Venue Staff cannot create online bookings; Customers cannot create walk-ins."),
        ("PRE-1: The caller is Staff in the UI, or Staff/Admin at API level.",
         "PRE-1: The caller is authenticated Venue Staff in both UI and API."),
        ("A1: API creation without payment\n1. An operator calls booking creation but does not call cash capture.\n2. The booking remains confirmed with zero paid and the full amount remaining.\nA2: Admin API operation\n1. An Admin may create the walk-in and record cash through protected APIs, although Admin checkout is intentionally hidden in React.",
         "A1: Pay later at the counter\n1. Venue Staff creates the walk-in without immediate capture.\n2. The booking remains pending with zero paid and the full amount remaining until Staff records enough cash to meet the deposit rule."),
        ("The current product decision says Staff handles the counter flow; Admin backend permission is broader than the React experience.",
         "Venue Staff alone handles the counter flow in both backend and React."),
        ("Both Customer Account and Staff workspace expose editors; Admin has backend authority but no matching editor.",
         "Customer Account and Venue Staff workspace expose editors; Admin is read-only for booking audit."),
        ("Staff is authenticated in the React workspace, or Staff/Admin calls the protected APIs.",
         "Venue Staff is authenticated in the React workspace and protected API."),
        ("EXC-2: Admin lacks the Staff calendar UI despite backend operator access.",
         "EXC-2: Admin and Customer callers are denied the Venue Staff operations calendar."),
        ("For manual confirmation, the caller is Staff/Admin.",
         "For manual confirmation, the caller is Venue Staff."),
        ("PRE-2: The caller is Staff/Admin.", "PRE-2: The caller is Venue Staff."),
        ("PRE-1: The caller is Staff/Admin.", "PRE-1: The caller is Venue Staff."),
        ("A1: Admin API rejection\n1. Admin submits the same protected status update through the API.\nA2: Slot reused later",
         "A1: Slot reused later"),
        ("The Customer owns it, or the caller is Staff/Admin.",
         "The Customer owns it, or the caller is Venue Staff."),
        ("A Customer owns it, or the caller is Staff/Admin.",
         "A Customer owns it, or the caller is Venue Staff."),
        ("Inactive-field and same-day elapsed-time checks are missing for a directly submitted new slot.",
         "Inactive fields, past/elapsed slots, blocked slots, and active booking conflicts are rejected."),
        ("Customer Account and Staff workspace expose the action; Admin has backend authority only.",
         "Customer Account and Venue Staff workspace expose the action; Admin is read-only for booking audit."),
        ("A1: Admin API check-in\n1. Admin may submit the protected status update without a corresponding Admin UI control.\nA2: Staff note",
         "A1: Staff note"),
        ("A2: Admin API action\n1. Admin may mark a confirmed booking no-show through the protected status API, although the Admin console has no lifecycle control.",
         "A2: Staff note\n1. Venue Staff may record the operational reason while marking no-show."),
        ("Customer cannot create walk-in bookings; Staff/Admin cannot create online bookings.",
         "Customers cannot create walk-in bookings; Venue Staff cannot create online bookings."),
        ("Admin has no checkout UI, although backend operator guards currently allow Admin walk-in/cash calls.",
         "Admin has no checkout UI or API authority for walk-in and counter-cash operations."),
        ("EXC-3: Unavailable/inactive service selections are rejected.",
         "EXC-3: Unavailable/inactive selections and edits at or after the booked start time are rejected."),
        ("Operator manual confirmation is allowed independently of paid amount; automatic capture requires paidAmount >= depositAmount.",
         "Both manual and automatic confirmation require paidAmount >= depositAmount."),
        ("The service does not enforce a check-in time window relative to slot start.",
         "Check-in opens 30 minutes before start and requires the configured deposit to be paid."),
        ("A2: Outstanding balance\n1. Completion still succeeds when remainingAmount is positive; the invoice preserves the outstanding balance.",
         "A2: Outstanding balance\n1. Completion is rejected until Venue Staff records the remaining counter payment."),
        ("A booking not checked in cannot be completed.",
         "A booking not checked in, not yet started, or carrying a remaining balance cannot be completed."),
        ("No arrival-time threshold is enforced.",
         "No-show is rejected until 15 minutes after the booked start time."),
    ]
    replace_text(document, replacements)


def align_sds_retake_rules(document: Document) -> None:
    """Keep the design signatures and role rules identical to the implemented API."""
    replace_text(
        document,
        [
            ("Exposes the staff/admin daily operations view.", "Exposes the Venue Staff daily operations view."),
            ("Exposes the Staff/Admin issue status endpoint.", "Exposes the Venue Staff issue status endpoint."),
            ("Allows Staff/Admin to create only walk-in bookings for an existing customer.",
             "Allows authenticated Venue Staff to create a walk-in booking for a matched Customer account or a first-time visitor contact."),
            ("Allows authenticated Venue Staff to create a walk-in booking for an existing Customer.",
             "Allows authenticated Venue Staff to create a walk-in booking for a matched Customer account or a first-time visitor contact."),
            ("Accepts deposit, full, or remaining for Staff/Admin counter payment but requires cash.",
             "Accepts deposit, full, or remaining for Venue Staff counter payment and requires cash."),
            ("Exposes Staff/Admin counter capture.", "Exposes Venue Staff counter capture."),
            ("Customer ownership and Staff/Admin access are enforced before the change; edits after check-in are blocked.",
             "Customer ownership or Venue Staff access is enforced before the change; edits after check-in are blocked."),
            ("Customer online checkout supports deposit/full through PayPal; Staff/Admin counter flow records cash and may use remaining.",
             "Customer online checkout supports deposit/full through PayPal; Venue Staff counter flow records cash and may use remaining."),
            ("capturePayPalOrder(Long bookingId, String orderId, ApiRequests.PayPalOrderCapture request)",
             "capturePayPalOrder(Long bookingId, String orderId)"),
            ("captureOrder(Long bookingId, String orderId, ApiRequests.PayPalOrderCapture request)",
             "captureOrder(Long bookingId, String orderId)"),
            (":reporterId", ":currentUserId"),
            (":assignedStaffId", ":currentStaffId"),
            (":staffId", ":currentStaffId"),
            ("USD 12 remains the explicit fallback when no band matches.",
             "a slot is unbookable when no active price band matches."),
        ],
    )


def set_update_fields(document: Document) -> None:
    settings = document.settings.element
    existing = settings.find(qn("w:updateFields"))
    if existing is None:
        existing = OxmlElement("w:updateFields")
        settings.append(existing)
    existing.set(qn("w:val"), "true")


def set_table_rows(table: Table, rows: list[list[str]]) -> None:
    while len(table.rows) < len(rows):
        table.add_row()
    while len(table.rows) > len(rows):
        table._tbl.remove(table.rows[-1]._tr)
    for row_index, values in enumerate(rows):
        tr_properties = table.rows[row_index]._tr.get_or_add_trPr()
        if tr_properties.find(qn("w:cantSplit")) is None:
            tr_properties.append(OxmlElement("w:cantSplit"))
        for column_index, value in enumerate(values):
            table.cell(row_index, column_index).text = value


def heading_level(element, document: Document) -> int | None:
    if not element.tag.endswith("}p"):
        return None
    paragraph = Paragraph(element, document)
    match = re.fullmatch(r"Heading (\d+)", paragraph.style.name or "")
    return int(match.group(1)) if match else None


def set_paragraph_text(element, document: Document, text: str, style: str | None = None) -> None:
    paragraph = Paragraph(element, document)
    paragraph.text = text
    if style:
        paragraph.style = style


def normalize_document_font(document: Document, font_name: str = "Arial") -> None:
    """Keep generated and template-retained text on one explicit typeface."""
    for style in document.styles:
        if getattr(style, "font", None) is None:
            continue
        style.font.name = font_name
        style_rpr = style.element.get_or_add_rPr()
        style_fonts = style_rpr.rFonts
        if style_fonts is None:
            style_fonts = OxmlElement("w:rFonts")
            style_rpr.insert(0, style_fonts)
        for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
            style_fonts.set(qn(f"w:{attribute}"), font_name)

    for part in document.part.package.parts:
        element = getattr(part, "element", None)
        if element is None:
            continue
        for run in element.xpath(".//w:r"):
            run_properties = run.get_or_add_rPr()
            run_fonts = run_properties.rFonts
            if run_fonts is None:
                run_fonts = OxmlElement("w:rFonts")
                run_properties.insert(0, run_fonts)
            for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
                run_fonts.set(qn(f"w:{attribute}"), font_name)


def renumber_uc_references_in_elements(elements: list) -> None:
    tokens = {old: f"[[OLD-UC-{old:02d}]]" for old in OLD_REFERENCE_TARGETS}
    for element in elements:
        for node in element.xpath(".//w:t"):
            text = node.text or ""
            for old, token in tokens.items():
                text = text.replace(f"UC-{old:02d}", token)
            for old, token in tokens.items():
                text = text.replace(token, OLD_REFERENCE_TARGETS[old])
            node.text = text


def renumber_rds_design_references(document: Document) -> None:
    elements = list(document.element.body.iterchildren())
    start = next(
        index for index, element in enumerate(elements)
        if heading_level(element, document) == 1
        and Paragraph(element, document).text.strip() == "III. Design Specifications"
    )
    renumber_uc_references_in_elements(elements[start:])


def table_set_label(table: Table, label: str, value: str) -> None:
    for row in table.rows:
        for index, cell in enumerate(row.cells[:-1]):
            if cell.text.strip() == label:
                row.cells[index + 1].text = value
                return


def update_rds_detail_table(table: Table, final_number: int, use_case: tuple) -> None:
    _, _, owner, actor, name, description = use_case
    uc_label = f"UC-{final_number:02d} {name}"
    for cell in table.rows[0].cells[1:]:
        cell.text = uc_label
    table_set_label(table, "Created By:", owner)
    table_set_label(table, "Primary Actor:", actor)
    table_set_label(table, "Description:", description)

    split_overrides = {
        7: {
            "Secondary Actors:": "GoalZone booking and issue repositories",
            "Trigger:": "An Admin opens People & Access and selects a Customer to inspect.",
            "Preconditions:": "PRE-1: The requester is authenticated as Admin.\nPRE-2: The selected account has the Customer role.",
            "Postconditions:": "POST-1: The Admin sees read-only Customer profile and operational activity context.\nPOST-2: No Customer data or access state is changed.",
            "Normal Flow:": "1. Admin opens the Customer directory.\n2. GoalZone returns Customer accounts only.\n3. Admin selects View.\n4. GoalZone loads bookings and issues visible as business context.\n5. The UI shows names, booking codes, fields, descriptions, statuses, and dates without exposing raw IDs as content.",
            "Alternative Flow:": "A1: No activity\n1. The account detail opens with an empty-state message.",
            "Exceptions:": "EXC-1: Non-Admin callers receive 403.\nEXC-2: A missing or non-Customer target is rejected.",
            "Other Information:": "UC-07 is deliberately read-only. The security mutation is isolated in UC-08 Lock/unlock customer account.",
        },
        8: {
            "Secondary Actors:": "VerificationEmailService and notification repository",
            "Trigger:": "An Admin chooses Lock or Unlock for a Customer account.",
            "Preconditions:": "PRE-1: The requester is authenticated as Admin.\nPRE-2: The target has the Customer role.\nPRE-3: Lock includes a non-blank reason.",
            "Postconditions:": "POST-1: status and lockReason reflect the decision.\nPOST-2: authVersion is incremented so older JWTs fail.\nPOST-3: The Customer receives the relevant email/in-app notification.",
            "Normal Flow:": "1. Admin selects a Customer and chooses Lock.\n2. The UI requires a reason.\n3. GoalZone validates the actor and target.\n4. It sets status=locked, stores the reason, increments authVersion, and notifies the Customer.\n5. Unlock later restores status=active, clears the reason, increments authVersion, and sends restoration notice.",
            "Alternative Flow:": "A1: Already locked/active\n1. GoalZone returns the current state without fabricating a second account state.",
            "Exceptions:": "EXC-1: Non-Admin callers receive 403.\nEXC-2: Blank lock reason or non-Customer target is rejected.",
            "Other Information:": "UC-08 changes access state and audit evidence; UC-07 remains the separate read-only directory/detail goal.",
        },
        13: {
            "Postconditions:": "POST-1: Missing non-overlapping slots for the requested date are materialized idempotently.\nPOST-2: Only matching future-start rows are displayed with price and availability; an available row may continue to checkout.",
            "Normal Flow:": "1. The visitor opens availability search.\n2. The visitor chooses a date and may filter by field type or field.\n3. GoalZone prepares any missing slots under the active venue schedule.\n4. GoalZone excludes elapsed starts and marks booked or blocked rows unavailable.\n5. GoalZone prices matching rows and displays the daily result.\n6. The visitor may select only an available row and continue to checkout.",
            "Exceptions:": "EXC-1: Past dates and stale direct checkout/create attempts are rejected.\nEXC-2: Dates beyond the configured generation horizon are rejected.\nEXC-3: A database or network error prevents results from refreshing.",
            "Other Information:": "Booked/blocked preservation and price matching are business rules within the search. Scheduled rolling-calendar extension is background behavior owned by UC-16, not an alternative visitor action.",
        },
        16: {
            "Secondary Actors:": "SlotGenerationService and Scheduler",
            "Trigger:": "Admin saves opening time, closing time, slot duration, or rolling horizon.",
            "Preconditions:": "PRE-1: The requester is authenticated as Admin.\nPRE-2: The combined rule set is complete.",
            "Postconditions:": "POST-1: The validated setting is stored.\nPOST-2: Only unbooked available auto-generated rows are rebuilt; bookings and blocks are preserved.",
            "Normal Flow:": "1. Admin opens Policies and edits a slot rule.\n2. GoalZone validates HH:mm times, duration, divisibility, and 1–90 day horizon.\n3. It stores the setting.\n4. SlotGenerationService rebuilds the rolling window idempotently.\n5. Search and operations calendar use the new rules.",
            "Alternative Flow:": "A1: Scheduled extension\n1. The scheduler extends the horizon without changing the configured rules.",
            "Exceptions:": "EXC-1: Opening must precede closing.\nEXC-2: Duration must be 30–360 minutes in 30-minute steps and divide the operating window.\nEXC-3: Horizon outside 1–90 days is rejected.",
            "Other Information:": "This Admin policy flow is separate from the Venue Staff exception flow in UC-17.",
        },
        17: {
            "Secondary Actors:": "Booking and slot repositories",
            "Trigger:": "Venue Staff must record or remove an operational field exception.",
            "Preconditions:": "PRE-1: The requester is authenticated as Staff.\nPRE-2: Field, future date, valid time range, and block reason are supplied.",
            "Postconditions:": "POST-1: Every overlapping eligible slot is blocked, or one blocked row is restored.\nPOST-2: Customer search and the operation calendar immediately reflect the exception.",
            "Normal Flow:": "1. Staff chooses field, date, time range, reason, and optional note.\n2. GoalZone loads overlapping slots and active bookings.\n3. With no conflict, it blocks every overlap or records a future exception row.\n4. Staff later selects Unblock to restore the chosen blocked row.\n5. The calendar refreshes.",
            "Alternative Flow:": "A1: Future exception\n1. Staff records the block before the date enters the generated horizon; later generation skips it.",
            "Exceptions:": "EXC-1: Past date, invalid range, or blank reason is rejected.\nEXC-2: Any overlapping active booking blocks the operation.\nEXC-3: Admin and Customer callers receive 403.",
            "Other Information:": "Block/Unblock is for exceptions only; UC-16 owns the ordinary automatic slot-generation policy.",
        },
        25: {
            "Secondary Actors:": "Booking, payment, and invoice repositories",
            "Trigger:": "Venue Staff receives a walk-in request from a registered Customer or a first-time visitor.",
            "Preconditions:": "PRE-1: The caller is authenticated Venue Staff in both UI and API.\nPRE-2: Staff either selects an active Customer matched by phone/email or supplies a first-time visitor's name and phone.\nPRE-3: A bookable future slot is selected; services and promotion are valid.",
            "Postconditions:": "POST-1: The slot is reserved under the Staff identity and exactly one customer identity: registered account or guest contact snapshot.\nPOST-2: Pay later remains pending; successful deposit/full cash confirms the booking and updates its invoice.\nPOST-3: A failed immediate cash transaction leaves neither a partial booking nor payment.",
            "Normal Flow:": "1. Staff opens Walk-in booking and searches by phone or email.\n2. GoalZone returns a matching active Customer and Staff confirms that account.\n3. Staff selects date, slot, services, and optional promotion.\n4. Staff chooses deposit or full cash payment.\n5. GoalZone validates the request and atomically stores the booking and cash transaction.\n6. GoalZone confirms the booking, updates the invoice, and opens it in Staff operations.",
            "Exceptions:": "EXC-1: Missing customer identity, unavailable slot, invalid service, or ineligible promotion is rejected.\nEXC-2: A first-time visitor requires a non-blank name and phone; email is optional.\nEXC-3: If immediate cash capture fails, the create-and-pay transaction is rolled back.\nEXC-4: Non-Staff callers receive 403.",
            "Other Information:": "The guest contact is stored only on the booking and does not create a Customer account. Guest walk-ins receive no membership discount, account notification, or PayPal checkout; Staff manages their booking and any cash refund at the venue.",
            "Assumptions:": "Staff verifies either the matched Customer account or the first-time visitor's contact details before confirming the booking.",
        },
    }
    for label, value in split_overrides.get(final_number, {}).items():
        table_set_label(table, label, value)

    behavior_overrides = {
        3: {
            "Postconditions:": "POST-1: The local session is cleared.\nPOST-2: authVersion is incremented so all JWTs issued earlier for the account are rejected.",
            "Other Information:": "Logout currently means sign out from all sessions because token revocation is account-version based rather than per-token storage.",
        },
        29: {
            "Normal Flow:": "1. Customer or Staff opens an eligible booking.\n2. The actor chooses a new date.\n3. GoalZone displays available slots for that date.\n4. The actor selects a destination and confirms.\n5. GoalZone revalidates availability, moves the booking, and recalculates invoice, balance, and refundable difference.",
        },
        31: {
            "Normal Flow:": "1. Staff selects an eligible confirmed booking at the venue.\n2. GoalZone checks the time window and minimum required payment.\n3. Staff reviews the confirmation and may add an audit note.\n4. GoalZone records checked-in status and the Staff identity.",
        },
        32: {
            "Normal Flow:": "1. Staff selects a checked-in booking after its start.\n2. GoalZone verifies that the remaining balance is zero.\n3. Staff reviews the confirmation and may add an audit note.\n4. GoalZone records completion and updates membership progress.",
        },
        33: {
            "Preconditions:": "PRE-1: Staff is authenticated.\nPRE-2: The booking is confirmed and at least 15 minutes past its start.\nPRE-3: Staff supplies a no-show note.",
            "Normal Flow:": "1. Staff selects the overdue confirmed booking.\n2. Staff chooses No-show and records what was checked.\n3. GoalZone asks for confirmation.\n4. GoalZone stores no-show status, note, time, and Staff identity without adding membership progress.",
        },
        35: {
            "Normal Flow:": "1. GoalZone shows deposit and full-payment choices for Customer checkout.\n2. For a walk-in, GoalZone also shows Pay later.\n3. The actor selects one option and reviews its paid and remaining amounts.\n4. The chosen option is passed to the corresponding online or counter-payment flow.",
        },
        37: {
            "Normal Flow:": "1. Staff opens a booking with an unpaid balance.\n2. Staff chooses Record remaining payment.\n3. GoalZone shows the exact cash amount and asks for confirmation.\n4. Staff confirms cash was received.\n5. GoalZone records the payment and reconciles the invoice.",
        },
        39: {
            "Normal Flow:": "1. The authorized user opens payment history.\n2. GoalZone filters transactions to the caller's authority.\n3. Each row shows payment code, booking code, field, booking date/time, method, paid time, status, and amount.\n4. Staff or Admin may open the related invoice context where authorized.",
        },
        41: {
            "Primary Actor:": "Staff",
            "Secondary Actors:": "PayPal Sandbox",
            "Preconditions:": "PRE-1: Venue Staff is authenticated.\nPRE-2: The refund request is approved and references a captured PayPal payment.",
            "Normal Flow:": "1. Staff opens an approved online refund case.\n2. Staff confirms provider processing.\n3. GoalZone sends the idempotent refund request to PayPal Sandbox.\n4. GoalZone accepts provider COMPLETED as final and stores provider evidence.\n5. GoalZone updates refundable totals and notifies the Customer.",
        },
        42: {
            "Primary Actor:": "Customer or Staff",
            "Secondary Actors:": "Admin (read-only audit), PayPal Sandbox",
            "Normal Flow:": "1. The booking owner opens an eligible booking.\n2. The owner enters a refund amount within the displayed maximum and a reason.\n3. GoalZone stores a requested refund.\n4. Staff reviews the request and records approve or reject with a note.\n5. Approved cash refunds may be completed by Staff; approved online refunds continue through UC-41.\n6. Admin may inspect the resulting audit trail without changing it.",
        },
        45: {
            "Normal Flow:": "1. The visitor opens active offers.\n2. GoalZone displays active campaigns and their eligibility conditions.\n3. The visitor chooses Use this code.\n4. GoalZone opens booking with that code filled in; final eligibility is checked in checkout.",
        },
        46: {
            "Normal Flow:": "1. Admin opens the promotion catalogue and chooses Add or Edit.\n2. Admin enters campaign dates, value, limits, conditions, status, and whether membership stacking is allowed.\n3. GoalZone validates the campaign.\n4. Admin saves it and GoalZone refreshes the catalogue only after success.",
        },
        48: {
            "Normal Flow:": "1. Customer opens Membership.\n2. GoalZone loads the current active tier and its qualification period.\n3. GoalZone shows the matching lifetime, monthly, or consecutive-week progress toward the next tier.\n4. The tier list marks unlocked levels by tier order, not by unrelated lifetime totals.",
        },
        54: {
            "Other Information:": "The range applies revenue by actual payment paidAt and completed refund processedAt. Booking commercial breakdown remains tied to the booking's slotDate.",
        },
        55: {
            "Other Information:": "Booking counts, lifecycle totals, peak time, and utilization use slotDate for the selected range.",
        },
    }
    for label, value in behavior_overrides.get(final_number, {}).items():
        table_set_label(table, label, value)
    table_set_label(table, "Alternative Flow:", ALTERNATIVE_FLOW_OVERRIDES[final_number])


def update_rds_catalogue(document: Document) -> None:
    header = ["ID", "Feature", "Use Case", "Use Case Description"]
    rows = [header]
    for number, use_case in enumerate(FINAL_USE_CASES, start=1):
        _, feature, _, _, name, description = use_case
        rows.append([f"UC-{number:02d}", feature, name, description])
    catalogues = [
        table for table in document.tables
        if table.rows and [cell.text.strip() for cell in table.rows[0].cells][:4] == header
    ]
    if len(catalogues) != 5:
        raise RuntimeError(f"Expected five RDS module catalogue tables, found {len(catalogues)}")
    for table, (_, first, last) in zip(catalogues, FINAL_MODULES):
        set_table_rows(table, [header] + rows[first:last + 1])


def rebuild_rds_requirement_sections(document: Document) -> None:
    body = document.element.body
    elements = list(body.iterchildren())
    section_start = next(
        index for index, element in enumerate(elements)
        if heading_level(element, document) == 1
        and Paragraph(element, document).text.strip() == "II. Requirement Specifications"
    )
    section_end = next(
        index for index, element in enumerate(elements)
        if index > section_start
        and heading_level(element, document) == 1
        and Paragraph(element, document).text.strip() == "III. Design Specifications"
    )

    source_blocks: dict[int, list] = {}
    module_samples: list = []
    index = section_start + 1
    while index < section_end:
        element = elements[index]
        level = heading_level(element, document)
        text = Paragraph(element, document).text.strip() if level else ""
        match = re.fullmatch(r"UC-(\d{2})\..+", text)
        if level == 2:
            module_samples.append(element)
            index += 1
            continue
        if level == 4 and match:
            end = index + 1
            while end < section_end:
                next_level = heading_level(elements[end], document)
                if next_level in {1, 2, 4}:
                    break
                end += 1
            source_blocks[int(match.group(1))] = elements[index:end]
            index = end
            continue
        index += 1

    if len(source_blocks) < 63 or len(module_samples) < 5:
        raise RuntimeError("Unexpected RDS requirement structure")

    for element in elements[section_start + 1:section_end]:
        body.remove(element)
    insert_at = body.index(elements[section_end])

    for module_index, (module_name, first, last) in enumerate(FINAL_MODULES):
        module = deepcopy(module_samples[module_index])
        set_paragraph_text(module, document, module_name)
        body.insert(insert_at, module)
        insert_at += 1
        for final_number in range(first, last + 1):
            use_case = FINAL_USE_CASES[final_number - 1]
            old_number = use_case[0]
            block = [deepcopy(element) for element in source_blocks[old_number]]
            renumber_uc_references_in_elements(block)
            set_paragraph_text(block[0], document, f"UC-{final_number:02d}. {use_case[4]}")
            detail = next((Table(element, document) for element in block if element.tag.endswith("}tbl")), None)
            if detail is None:
                raise RuntimeError(f"RDS detail table missing for old UC-{old_number:02d}")
            update_rds_detail_table(detail, final_number, use_case)
            for element in block:
                body.insert(insert_at, element)
                insert_at += 1


def remove_rows_matching(table: Table, column: int, value: str) -> None:
    for row in list(table.rows)[1:]:
        if row.cells[column].text.strip() == value:
            table._tbl.remove(row._tr)


def remove_heading_subtree(document: Document, heading_text: str, heading: int) -> None:
    body = document.element.body
    elements = list(body.iterchildren())
    start = next(
        (index for index, element in enumerate(elements)
         if heading_level(element, document) == heading
         and Paragraph(element, document).text.strip() == heading_text),
        None,
    )
    if start is None:
        return
    end = start + 1
    while end < len(elements):
        level = heading_level(elements[end], document)
        if level is not None and level <= heading:
            break
        if elements[end].tag.endswith("}sectPr"):
            break
        end += 1
    for element in elements[start:end]:
        body.remove(element)


def align_rds_screens(document: Document) -> None:
    for table in document.tables:
        if not table.rows:
            continue
        header = [cell.text.strip() for cell in table.rows[0].cells]
        if header[:4] == ["#", "Feature", "Screen", "Description"]:
            remove_rows_matching(table, 2, "Availability Assistant")
            for number, row in enumerate(table.rows[1:], start=1):
                row.cells[0].text = str(number)
        elif header[:5] == ["Screen", "Guest", "Customer", "Staff", "Admin"]:
            remove_rows_matching(table, 0, "Availability Assistant")
    remove_heading_subtree(document, "j. Availability Assistant", 4)


def normalized_sequence_filename(number: int, name: str) -> str:
    slug = name.lower().replace("/", "-").replace(" ", "-")
    slug = re.sub(r"[^a-z0-9-]+", "", slug)
    slug = re.sub(r"-+", "-", slug).strip("-")
    special = {
        1: "register-account",
        6: "recover-password",
        8: "lock-unlock-customer-account",
        13: "search-available-fields-by-date-time",
        31: "check-in-booking",
        36: "pay-online-via-paypal-sandbox",
    }
    return f"UC-{number:02d}-{special.get(number, slug)}.png"


def rebuild_sds_use_case_sections(document: Document) -> None:
    body = document.element.body
    elements = list(body.iterchildren())
    section_start = next(
        index for index, element in enumerate(elements)
        if heading_level(element, document) == 1
        and Paragraph(element, document).text.strip() == "II. Code Designs"
    )
    section_end = next(
        (index for index, element in enumerate(elements[section_start + 1:], start=section_start + 1)
         if heading_level(element, document) == 1),
        next(index for index, element in enumerate(elements) if element.tag.endswith("}sectPr")),
    )
    source_blocks: dict[int, list] = {}
    index = section_start + 1
    while index < section_end:
        element = elements[index]
        if heading_level(element, document) == 2:
            text = Paragraph(element, document).text.strip()
            match = re.fullmatch(r"(\d+)\..+", text)
            if match:
                end = index + 1
                while end < section_end and heading_level(elements[end], document) != 2:
                    end += 1
                source_blocks[int(match.group(1))] = elements[index:end]
                index = end
                continue
        index += 1
    if len(source_blocks) < 63:
        raise RuntimeError("Unexpected SDS code-design structure")

    for element in elements[section_start + 1:section_end]:
        body.remove(element)
    insert_at = body.index(elements[section_end])

    for final_number, use_case in enumerate(FINAL_USE_CASES, start=1):
        old_number, _, _, _, name, _ = use_case
        source = source_blocks[old_number]
        if final_number == 7:
            cutoff = next(
                index for index, element in enumerate(source)
                if element.tag.endswith("}p")
                and Paragraph(element, document).text.strip() == "e. Account lock/unlock design"
            )
            block = [deepcopy(element) for element in source[:cutoff]]
        elif final_number == 8:
            start = next(
                index for index, element in enumerate(source)
                if element.tag.endswith("}p")
                and Paragraph(element, document).text.strip() == "e. Account lock/unlock design"
            )
            block = [deepcopy(source[0])] + [deepcopy(element) for element in source[start + 1:]]
        else:
            block = [deepcopy(element) for element in source]

        set_paragraph_text(block[0], document, f"{final_number}. {name}")
        for element in block[1:]:
            if not element.tag.endswith("}p"):
                continue
            paragraph = Paragraph(element, document)
            text = paragraph.text.strip()
            if final_number == 8 and text.startswith("e."):
                replacements = {"e.1 Class Diagram": "a. Class Diagram", "e.2 Class Specifications": "b. Class Specifications", "e.3 Sequence Diagram": "c. Sequence Diagram", "e.4 Database Queries": "d. Database Queries"}
                if text in replacements:
                    paragraph.text = replacements[text]
                    paragraph.style = "Heading 3"
            elif text.startswith("a. Class Diagram"):
                paragraph.text = "a. Class Diagram"
            elif text.startswith("c. Sequence Diagram"):
                paragraph.text = "c. Sequence Diagram"
            elif re.fullmatch(r"== .+ ==", text):
                paragraph.text = f"== {name} =="

        old_id = f"UC-{old_number:02d}"
        new_id = f"UC-{final_number:02d}"
        for element in block:
            if element.tag.endswith("}p"):
                replace_across_runs(Paragraph(element, document), old_id, new_id)
            elif element.tag.endswith("}tbl"):
                table = Table(element, document)
                for row in table.rows:
                    for cell in row.cells:
                        for paragraph in cell.paragraphs:
                            replace_across_runs(paragraph, old_id, new_id)
        for element in block:
            body.insert(insert_at, element)
            insert_at += 1


def replace_all_sds_diagrams(document: Document) -> None:
    for number, (use_case, class_image) in enumerate(zip(FINAL_USE_CASES, CLASS_IMAGE_NAMES), start=1):
        heading = f"{number}. {use_case[4]}"
        replace_drawing_after(document, heading, "a. Class Diagram", CLASS_RENDER / class_image)
        replace_drawing_after(
            document,
            heading,
            "c. Sequence Diagram",
            SEQUENCE_RENDER / normalized_sequence_filename(number, use_case[4]),
        )


def update_package_tables(document: Document) -> None:
    descriptions = {
        "com.swp391.backend.controller": (
            "8 Java file(s): AccountController, ApiExceptionHandler, "
            "BookingController, FieldOperationController, "
            "ImageUploadController, PaymentController, PromotionReportController, and SystemController."
        ),
        "com.swp391.backend.service": (
            "15 Java file(s): AccountService, ApiException, "
            "BookingMaintenanceScheduler, BookingWorkflowService, DomainSupportService, "
            "FieldOperationService, ImageStorageService, PayPalCheckoutService, "
            "PaymentWorkflowService, PromotionReportService, SlotGenerationScheduler, "
            "SlotGenerationService, SystemQueryService, "
            "VerificationEmailDelivery, and VerificationEmailService."
        ),
    }
    for table in document.tables:
        if len(table.columns) != 3 or not table.rows:
            continue
        header = [cell.text.strip() for cell in table.rows[0].cells]
        if header[:3] != ["No", "Package", "Description"]:
            continue
        for row in table.rows[1:]:
            package = row.cells[1].text.strip()
            if package in descriptions:
                row.cells[2].text = descriptions[package]


def update_database_table_descriptions(document: Document) -> None:
    for table in document.tables:
        if len(table.columns) != 3 or not table.rows:
            continue
        header = [cell.text.strip() for cell in table.rows[0].cells]
        if header[:3] != ["No", "Table", "Description"]:
            continue
        for row in table.rows[1:]:
            name = row.cells[1].text.strip()
            if name == "app_user":
                row.cells[2].text = (
                    "Primary key: user_id. Stores Customer, Staff, and Admin identities. "
                    "status is the single persisted source of truth for active/inactive/locked access; "
                    "lock_reason explains an administrative Customer lock and auth_version revokes JWTs."
                )
            elif name == "field_type":
                row.cells[2].text = (
                    "Primary key: field_type_id. Normalizes reusable field-type name and player capacity "
                    "so multiple fields do not duplicate the same catalogue metadata."
                )
            elif name == "field_price":
                row.cells[2].text = (
                    "Primary key: field_price_id. Keeps effective weekday/weekend and time-band prices "
                    "separate from field identity so pricing can change over time without rewriting a field."
                )
            elif name == "booking":
                row.cells[2].text = (
                    "Primary key: booking_id. Reserves one slot and stores pricing/lifecycle state. "
                    "Identity is exclusive: customer_id for a registered Customer, or guest_name and guest_phone "
                    "(guest_email optional) for a walk-in visitor. Guest details are a booking snapshot, not an account."
                )
            elif name == "booking_promotion":
                row.cells[2].text = (
                    "Primary key: booking_promotion_id. Snapshots the applied code and discount for one booking. "
                    "usage_counted records whether the pending/active booking currently reserves campaign capacity."
                )
            elif name == "system_setting":
                row.cells[2].text = (
                    "Primary key: setting_id. Stores administrator-editable deposit, refund, payment-timeout, "
                    "reminder, and automatic slot-generation policies as key/value configuration; "
                    "it is not redundant field data."
                )


def update_slot_use_case_descriptions(document: Document) -> None:
    catalogue = {
        "UC-12": (
            "Generate missing bookable slots from active Admin rules, then search the requested date and optional "
            "field type, calculate current field price, and label each row as available, booked, or blocked."
        ),
        "UC-15": (
            "Admin configures opening time, closing time, slot duration, and rolling horizon. Staff preserves "
            "maintenance, private-event, VIP/internal-use, and other exceptions through Block/Unblock."
        ),
        "UC-16": (
            "Ensure the selected date is materialized from Admin slot rules, then display every field slot with "
            "available, blocked, or booking-derived operational status."
        ),
    }
    detailed = {
        "UC-12": {
            "Primary Actor:": "Guest or Customer",
            "Description:": catalogue["UC-12"],
            "Preconditions:": (
                "PRE-1: At least one active field exists and the four slot-generation rules are active.\n"
                "PRE-2: No authentication is required for ordinary slot search."
            ),
            "Postconditions:": (
                "POST-1: Missing non-overlapping slots for the requested date are materialized idempotently.\n"
                "POST-2: Matching rows are displayed with price and availability; an available row may continue to checkout."
            ),
            "Normal Flow:": (
                "A. Search available fields by date/time\n"
                "1. The visitor chooses a date and optional field type/field.\n"
                "2. The frontend requests GET /api/slots/search.\n"
                "3. The backend reads opening time, closing time, duration, and generation horizon.\n"
                "4. Missing slots are generated for every active field without overlapping blocked/booked rows.\n"
                "5. The service derives booking state and price, filters by field type, and returns the daily list.\n"
                "6. The UI enables only rows with available=true."
            ),
            "Alternative Flow:": (
                "A1: No date supplied\n1. The service defaults to tomorrow and generates it on demand.\n"
                "A2: Existing exception\n1. A booked or blocked row is preserved and returned as unavailable.\n"
                "A3: Daily scheduler\n1. The midnight job extends the rolling calendar before any customer search."
            ),
            "Exceptions:": (
                "EXC-1: Past dates are rejected.\n"
                "EXC-2: Dates beyond the configured generation horizon are rejected.\n"
                "EXC-3: Invalid or non-divisible schedule rules are rejected when Admin saves them.\n"
                "EXC-4: A database or network error prevents results from refreshing."
            ),
            "Other Information:": (
                "Implementation rule: generation is idempotent and synchronized in the application instance; "
                "the database unique/exclusion constraints remain the final overlap guard."
            ),
            "Assumptions:": (
                "The venue uses one global operating schedule for all active fields; field-specific hours may be "
                "introduced later without changing the booking lifecycle."
            ),
        },
        "UC-15": {
            "Primary Actor:": "Admin and Staff",
            "Secondary Actors:": "GoalZone slot-generation scheduler and booking database",
            "Trigger:": (
                "Admin changes automatic availability rules, or Staff blocks/unblocks an exceptional field period."
            ),
            "Description:": catalogue["UC-15"],
            "Preconditions:": (
                "PRE-1: Admin is authenticated to edit slot-generation rules, or Staff is authenticated to manage exceptions.\n"
                "PRE-2: The target field exists; Block requires date, start, end, and reason."
            ),
            "Postconditions:": (
                "POST-1: A valid rule change rebuilds unbooked auto-generated availability while preserving bookings and blocks.\n"
                "POST-2: Block affects every overlapping slot without an active booking; Unblock restores one blocked row.\n"
                "POST-3: Customer search immediately reflects the resulting calendar."
            ),
            "Normal Flow:": (
                "A. Configure automatic availability\n"
                "1. Admin opens Policies and edits opening time, closing time, duration, or horizon.\n"
                "2. The backend validates the combined rule set.\n"
                "3. Unbooked auto-generated rows in the rolling window are rebuilt; booked and blocked rows are preserved.\n"
                "B. Block an exception\n"
                "1. Staff opens Field schedule and chooses a field, date, time range, and reason.\n"
                "2. The backend rejects the request if any overlapping slot has an active booking.\n"
                "3. Every overlapping slot is marked blocked; if none exists, a future exception row is created.\n"
                "4. Calendar and public availability refresh."
            ),
            "Alternative Flow:": (
                "A1: Unblock exception\n1. Staff selects Unblock and the backend clears reason/note and restores availability.\n"
                "A2: Future private event\n1. Staff may create a blocked exception before that date enters the rolling horizon; "
                "later generation skips its range."
            ),
            "Exceptions:": (
                "EXC-1: Rule time must use HH:mm; opening must precede closing.\n"
                "EXC-2: Duration must be 30-360 minutes in 30-minute steps and divide the operating window exactly.\n"
                "EXC-3: Horizon must be 1-90 days.\n"
                "EXC-4: Past blocks, missing reasons, invalid ranges, and ranges containing active bookings are rejected."
            ),
            "Other Information:": (
                "Block/Unblock is intentionally retained for maintenance, weather, special events, VIP/internal use, "
                "and other operational exceptions. Manual Add Slot is not part of the normal workflow."
            ),
            "Assumptions:": (
                "Only unbooked rows originally generated by the system are replaceable when rules change."
            ),
        },
        "UC-16": {
            "Description:": catalogue["UC-16"],
            "Preconditions:": (
                "PRE-1: The caller has Staff or Admin backend authority; Staff owns the current React calendar.\n"
                "PRE-2: Slot-generation rules are valid for the requested date."
            ),
            "Postconditions:": (
                "POST-1: The requested date exists in the rolling calendar and the operator sees available, blocked, and booked states.\n"
                "POST-2: Viewing does not change bookings or block exceptions."
            ),
            "Normal Flow:": (
                "A. View field operation calendar\n"
                "1. Staff opens Field schedule or changes its date.\n"
                "2. The frontend requests GET /api/operations/calendar?date=... .\n"
                "3. The backend materializes any missing slots from Admin rules.\n"
                "4. It loads slots and bookings, sorts by field/time, and derives operationalStatus.\n"
                "5. The UI displays field, time, booking/customer or block reason, and status."
            ),
            "Alternative Flow:": (
                "A1: Date omitted\n1. The backend defaults to today and generates remaining non-elapsed slots.\n"
                "A2: No active fields\n1. The calendar returns an empty list."
            ),
            "Exceptions:": (
                "EXC-1: Unauthenticated and Customer callers are forbidden.\n"
                "EXC-2: Past or beyond-horizon operation dates are rejected by the generation rules."
            ),
            "Other Information:": (
                "Implementation rule: booking status takes precedence over slot status in the operational display; "
                "Block/Unblock remains adjacent to this calendar."
            ),
            "Assumptions:": "The calendar is a daily operations view rather than a month-grid planning engine.",
        },
    }
    for table in document.tables:
        if not table.rows:
            continue
        header = [cell.text.strip() for cell in table.rows[0].cells]
        if header[:4] == ["ID", "Feature", "Use Case", "Use Case Description"]:
            for row in table.rows[1:]:
                uc_id = row.cells[0].text.strip()
                if uc_id in catalogue:
                    row.cells[3].text = catalogue[uc_id]
            continue
        first_row = " | ".join(cell.text.strip() for cell in table.rows[0].cells)
        uc_id = next((candidate for candidate in detailed if candidate in first_row), None)
        if uc_id is None:
            continue
        values = detailed[uc_id]
        for label, value in values.items():
            table_set_label(table, label, value)


def update_rds() -> None:
    document = Document(RDS_INPUT)
    update_cover_metadata(document)
    renumber_rds_design_references(document)
    replacements = [
        (
            "A lock stores status `locked`, `account_locked = true`, and the required reason.",
            "A lock stores status `locked` as the single persisted access state together with the required reason.",
        ),
        ("status/account_locked/lock_reason/auth_version", "status/lock_reason/auth_version"),
        ("`accountLocked` is false", "`status` is `active`"),
        (
            "Rule-based discovery page that ranks live available slots and continues a selected result into real booking.",
            "Availability search page that shows generated, priced, non-blocked field slots and continues a selected result into checkout.",
        ),
        (
            "UC-62 and UC-63 are implemented as a single rule-based discovery-to-booking journey while retaining separate traceability identifiers.",
            "Availability discovery is covered by UC-13 Search available fields by date/time; duplicate assistant and ranking journeys are outside the retake scope.",
        ),
        (
            "UC-62/63 use an explainable rule-based assistant over live data; external LLM/RAG integration remains optional and is not required for classroom delivery.",
            "GoalZone uses deterministic availability, pricing, promotion, and membership rules over its own database.",
        ),
        ("Ask smart assistant for available fields", "Ask availability assistant"),
        ("payment is captured or staff approval is recorded.", "the required payment is captured."),
        ("Booking is confirmed after valid payment or staff approval.",
         "Booking is confirmed after valid required payment; there is no manual approval queue."),
        ("Pending booking request was rejected by staff.",
         "Historical compatibility status; current invalid requests fail during booking/payment instead of Staff review."),
        ("UC-13 is the ordinary public date/type search. UC-13/UC-13 add explainable ranking by preferred time, budget and eligible promotions.",
         "UC-13 is the single ordinary public date/time/type availability search over generated and priced live slots."),
        ("USD 12 remains the explicit fallback when no band matches.",
         "A slot is not bookable when no active effective price band matches."),
        ("Admin has no checkout UI even where operator APIs permit Admin.",
         "Admin has no checkout UI or API authority; Venue Staff owns walk-in cash operations."),
        ("Booked-value report metrics exclude rejected/expired bookings but currently include cancelled/no-show snapshots.",
         "Booked commercial value excludes cancelled/rejected/expired rows; utilization counts only qualifying operational statuses."),
        ("UC-29/UC-30 composes calendar, issue, reschedule, cancel, refund and notification functions; there is no dedicated conflict entity.",
         "Booking conflict is an alternate flow that composes issue, UC-29 reschedule, UC-30 cancel, refund, and notification behavior; there is no dedicated conflict entity."),
        ("UC-13/UC-13 are deterministic and explainable; no external LLM, RAG knowledge base or fabricated result source is included.",
         "The retired assistant/ranking journey is outside scope; UC-13 deterministic live availability search is authoritative."),
        ("Each selected service passes UC-20 availability checks.",
         "Each selected service passes the included active/quantity/overlapping-stock validation."),
        ("Any UC-20 validation failure rejects the preview/create", "Any included service-validation failure rejects the preview/create"),
        ("Successful capture delegates to UC-36 state reconciliation.",
         "Successful verified capture is an included step that reconciles booking, payment, invoice, and confirmation state."),
        ("Admin also has a per-Customer activity modal elsewhere, but UC-56", "UC-10 provides a per-Customer operational modal, while UC-56"),
        ("Flyway V11", "Flyway V20"),
        ("Flyway V15", "Flyway V20"),
        ("Flyway V16", "Flyway V20"),
        ("Flyway V17", "Flyway V20"),
        ("Flyway V18", "Flyway V20"),
    ]
    replace_text(document, replacements)
    replace_xml_attribute_text(
        document,
        [
            ("Flyway V1 through V11", "Flyway V1 through V20"),
            ("Flyway V1 through V15", "Flyway V1 through V20"),
            ("Flyway V1 through V16", "Flyway V1 through V20"),
            ("Flyway V1 through V17", "Flyway V1 through V20"),
            ("Flyway V1 through V18", "Flyway V1 through V20"),
            ("UC-64", "UC-63"),
        ],
    )
    update_slot_use_case_descriptions(document)
    align_rds_retake_rules(document)
    remove_rds_external_diagram_section(
        document,
        "a. Use Case Diagrams",
        "b. Descriptions",
    )

    update_rds_catalogue(document)
    rebuild_rds_requirement_sections(document)

    remove_rds_external_diagram_section(
        document,
        "2.1 Screens Flow",
        "2.2 Screen Descriptions",
    )

    # Screen description and authorization are corrected against actual routes/roles.
    screen_descriptions = {
        "Unavailable Slot Management": (
            "Venue Staff reviews the daily field schedule and blocks or unblocks operational slot exceptions. "
            "Admin slot-generation rules are configured separately in the policy workspace."
        ),
        "Issue Report / Issue Management": (
            "Customer or Staff reports a field, booking, or service issue; Venue Staff resolves or rejects it; "
            "Admin has read-only audit access."
        ),
        "Walk-in Booking": (
            "Venue Staff creates a walk-in booking for either a Customer matched by phone/email or a first-time "
            "visitor identified by name and phone, then records deposit/full cash or leaves payment pending."
        ),
        "Booking Detail": (
            "Customer views an owned booking, Venue Staff performs eligible lifecycle actions, and Admin reviews "
            "booking and billing data in read-only audit mode."
        ),
        "Booking Operations": (
            "Venue Staff uses the booking queue and selected-booking controls for reschedule, cancellation, "
            "check-in, completion, no-show, services, cash payment, invoice, and transaction history."
        ),
        "Payment History": (
            "Customer views owned transactions; Venue Staff and Admin review role-authorized payment records, "
            "with Admin access remaining read-only."
        ),
        "Invoice Detail": (
            "Displays the generated invoice and settlement state to the booking owner or Venue Staff; Admin access "
            "is read-only audit."
        ),
        "Refund Management": (
            "Customer submits an eligible refund request, Venue Staff reviews and processes it, and Admin inspects "
            "the refund audit trail without changing it."
        ),
        "Policy Management": (
            "Admin policy workspace for deposit, cancellation/refund, payment timeout, reminder, and automatic "
            "slot opening/closing/duration/horizon rules."
        ),
    }
    for table in document.tables:
        if not table.rows:
            continue
        header = [cell.text.strip() for cell in table.rows[0].cells]
        if header[:4] == ["#", "Feature", "Screen", "Description"]:
            for row in table.rows[1:]:
                screen_name = row.cells[2].text.strip()
                if screen_name == "Booking Calendar":
                    screen_name = "Booking Operations"
                    row.cells[2].text = screen_name
                elif screen_name == "Booking and Refund Policies":
                    screen_name = "Policy Management"
                    row.cells[2].text = screen_name
                if screen_name in screen_descriptions:
                    row.cells[3].text = screen_descriptions[screen_name]
        if header[:5] == ["Screen", "Guest", "Customer", "Staff", "Admin"]:
            access = {
                "Landing Page": {"Guest", "Customer"},
                "Login Screen": {"Guest", "Customer", "Staff", "Admin"},
                "Register Account": {"Guest"},
                "Forgot/Reset Password": {"Guest"},
                "Profile / Edit Profile": {"Customer", "Staff", "Admin"},
                "Customer Management": {"Admin"},
                "Staff Management": {"Admin"},
                "Field List": {"Guest", "Customer", "Staff", "Admin"},
                "Field Detail": {"Guest", "Customer", "Staff", "Admin"},
                "Search Available Fields": {"Guest", "Customer", "Staff"},
                "Field Management": {"Admin"},
                "Field Pricing Management": {"Admin"},
                "Extra Service Selection": {"Guest", "Customer", "Staff"},
                "Extra Service Management": {"Admin"},
                "Unavailable Slot Management": {"Staff"},
                "Issue Report / Issue Management": {"Customer", "Staff", "Admin"},
                "Online Booking": {"Customer"},
                "Walk-in Booking": {"Staff"},
                "Booking Detail": {"Customer", "Staff", "Admin"},
                "My Bookings": {"Customer"},
                "Booking Operations": {"Staff"},
                "Reschedule Booking": {"Customer", "Staff"},
                "Cancellation Preview": {"Customer", "Staff"},
                "Checkout Summary": {"Guest", "Customer", "Staff"},
                "Payment Sandbox": {"Customer"},
                "Payment History": {"Customer", "Staff", "Admin"},
                "Invoice Detail": {"Customer", "Staff", "Admin"},
                "Refund Management": {"Customer", "Staff", "Admin"},
                "Policy Management": {"Admin"},
                "Promotion List": {"Guest", "Customer", "Staff", "Admin"},
                "Promotion Management": {"Admin"},
                "Membership Benefits": {"Guest", "Customer", "Staff", "Admin"},
                "Membership Progress": {"Customer"},
                "Membership Rule Management": {"Admin"},
                "Notification List": {"Customer", "Staff", "Admin"},
                "Revenue Report": {"Admin"},
                "Booking Report": {"Admin"},
                "Customer Account state Report": {"Admin"},
                "Customer Account State Report": {"Admin"},
            }
            for row in table.rows[1:]:
                screen_name = row.cells[0].text.strip()
                if screen_name == "Booking Calendar":
                    screen_name = "Booking Operations"
                    row.cells[0].text = screen_name
                elif screen_name == "Booking and Refund Policies":
                    screen_name = "Policy Management"
                    row.cells[0].text = screen_name
                allowed = access.get(screen_name, set())
                for column, role in enumerate(["Guest", "Customer", "Staff", "Admin"], start=1):
                    row.cells[column].text = "X" if role in allowed else ""

    align_rds_screens(document)
    replace_text(
        document,
        [
            ("UC-07/UC-08, UC-10", "UC-07, UC-08, UC-10"),
            ("UC-16/UC-17, UC-18", "UC-17, UC-18"),
            ("Related use cases: UC-16, UC-17, UC-18", "Related use cases: UC-17, UC-18"),
            ("Detailed as-built SQL is specified under UC-16, UC-17, UC-18 in the SDS Database Queries sections.",
             "Detailed as-built SQL is specified under UC-17, UC-18 in the SDS Database Queries sections."),
            ("Staff blocks/unblocks field time with a reason and reviews daily field operations.",
             "Venue Staff reviews daily field operations and blocks or unblocks operational slot exceptions; Admin configures ordinary slot generation separately."),
            ("Customer/Staff reports a field, booking, or service issue; Staff resolves it and notifies the reporter.",
             "Customer or Staff reports a field, booking, or service issue; Venue Staff resolves or rejects it and notifies the reporter; Admin reviews the history in read-only mode."),
            ("Role/ownership-protected view of booking, billing, and allowed lifecycle actions.",
             "Customer views an owned booking; Venue Staff performs eligible lifecycle actions; Admin reviews booking and billing data in read-only audit mode."),
            ("Shows real persisted transaction history subject to booking ownership/operator access.",
             "Shows persisted transaction history subject to ownership or operational access; Admin access is read-only audit."),
            ("Displays the generated booking invoice and current settlement state.",
             "Displays the generated booking invoice and current settlement state; Admin access is read-only audit."),
            ("Customer submits an eligible request; Staff reviews it and completes PayPal or cash refund workflows.",
             "Customer submits an eligible request; Venue Staff reviews and completes PayPal or cash refund workflows; Admin reviews the audit trail without changing it."),
            ("Admin maintains deposit, payment-timeout, cancellation/refund, and reminder settings.",
             "Admin maintains slot-generation, deposit, payment-timeout, cancellation/refund, and reminder settings."),
            ("f. Booking and Refund Policies", "f. Policy Management"),
            ("Related use cases: UC-43, UC-44", "Related use cases: UC-16, UC-43, UC-44"),
            ("Detailed as-built SQL is specified under UC-43, UC-44 in the SDS Database Queries sections.",
             "Detailed as-built SQL is specified under UC-16, UC-43, UC-44 in the SDS Database Queries sections."),
            ("UC-20, UC-20, UC-21", "UC-20, UC-21"),
            ("UC-25, UC-36, UC-37", "UC-25, UC-35, UC-37"),
            ("UC-26, UC-36, UC-24, UC-31, UC-32, UC-33", "UC-26, UC-31, UC-32, UC-33, UC-37, UC-40"),
            ("UC-28, UC-36, UC-24, UC-31, UC-32, UC-33, UC-29/UC-30", "UC-28, UC-29, UC-30, UC-31, UC-32, UC-33, UC-37"),
            ("UC-29, UC-29/UC-30", "UC-29"),
            ("UC-30, UC-30", "UC-30"),
            ("UC-36, UC-36, UC-38", "UC-36, UC-38"),
            ("UC-40, UC-40", "UC-40"),
            ("Customer Account state Report", "Customer Account State Report"),
            ("Each selected service passes UC-20 availability checks.",
             "Each selected service passes included active, quantity, and overlapping-stock validation."),
            ("Any UC-20 validation failure rejects the preview/create",
             "Any included service-validation failure rejects the preview/create"),
            ("Successful capture delegates to UC-36 state reconciliation.",
             "Successful verified capture is an included step that reconciles booking, payment, invoice, and confirmation state."),
            ("Admin also has a per-Customer activity modal elsewhere, but UC-56",
             "UC-10 provides a per-Customer operational modal, while UC-56"),
            ("UC-13 and UC-13 are implemented as a single rule-based discovery-to-booking journey while retaining separate traceability identifiers.",
             "UC-13 is the single deterministic availability-search journey over generated, priced GoalZone slots."),
            ("UC-13/63 use an explainable rule-based assistant over live data; external LLM/RAG integration remains optional and is not required for classroom delivery.",
             "The assistant/ranking journey is removed; UC-13 deterministic live availability search is sufficient for the classroom scope."),
            ("UC-13/63 are deterministic and explainable; no external LLM, RAG knowledge base or fabricated result source is included.",
             "UC-13 deterministically reads live GoalZone availability; no assistant, LLM, or fabricated result source is included."),
            ("UC-13 is the ordinary public date/type search. UC-13/63 add explainable ranking by preferred time, budget and eligible promotions.",
             "UC-13 is the single public date/time/type availability search over generated, priced live slots."),
            ("Assistant scope", "Availability scope"),
        ],
    )
    update_package_tables(document)
    update_database_table_descriptions(document)
    replace_text(document, [("Flyway V20", "Flyway V21"), ("through V20", "through V21")])
    replace_xml_attribute_text(document, [("Flyway V20", "Flyway V21"), ("through V20", "through V21")])
    set_update_fields(document)
    normalize_document_font(document)
    OUTPUT.mkdir(parents=True, exist_ok=True)
    document.save(RDS_OUTPUT)


def insert_spec_after(
    document: Document,
    anchor: Table,
    name: str,
    rows: list[list[str]],
) -> Table:
    heading_xml = deepcopy(anchor._tbl.getprevious())
    heading = Paragraph(heading_xml, document)
    heading.clear()
    heading.add_run(name).bold = True
    heading.paragraph_format.space_before = Inches(0.16)
    heading.paragraph_format.space_after = Inches(0.05)
    heading.paragraph_format.keep_with_next = True
    clone = deepcopy(anchor._tbl)
    anchor._tbl.addnext(heading_xml)
    heading._p.addnext(clone)
    table = Table(clone, document)
    set_table_rows(table, rows)
    return table


def update_sds_slot_generation_specs(document: Document) -> None:
    sections = {
        "12. Search available fields by date/time": {
            "FieldOperationService": [
                ["No", "Method", "Description"],
                [
                    "01",
                    "searchSlots(date : LocalDate, fieldTypeId : Long) : List<Map<String, Object>>",
                    "Ensures the requested day is materialized, then returns active-field slots with booking state, availability, and calculated USD price.",
                ],
            ],
            "SlotRepository": [
                ["No", "Method", "Description"],
                [
                    "01",
                    "findBySlotDate(date : LocalDate) : List<Slot>",
                    "Loads the materialized slots for the requested date.",
                ],
                [
                    "02",
                    "findByField_FieldIdAndSlotDate(fieldId : Long, date : LocalDate) : List<Slot>",
                    "Supports overlap-safe materialization for one field and date.",
                ],
            ],
        },
        "15. Manage unavailable slots": {
            "FieldOperationService": [
                ["No", "Method", "Description"],
                [
                    "01",
                    "blockSlot(request : ApiRequests.SlotBlock) : Map<String, Object>",
                    "Blocks every overlapping generated slot, rejects active-booking conflicts, or records a future manual exception outside the current horizon.",
                ],
                [
                    "02",
                    "unblockSlot(slotId : Long) : Map<String, Object>",
                    "Returns a blocked exception to available without recreating the operating calendar manually.",
                ],
            ],
            "SlotRepository": [
                ["No", "Method", "Description"],
                [
                    "01",
                    "findByField_FieldIdAndSlotDate(fieldId : Long, date : LocalDate) : List<Slot>",
                    "Loads every same-day row so Block can cover all overlapping generated periods.",
                ],
                [
                    "02",
                    "findBySlotDateBetweenOrderBySlotDateAscStartTimeAsc(from : LocalDate, to : LocalDate) : List<Slot>",
                    "Loads the rolling window for a safe Admin-rule rebuild.",
                ],
            ],
        },
        "16. View field operation calendar": {
            "FieldOperationService": [
                ["No", "Method", "Description"],
                [
                    "01",
                    "operationCalendar(date : LocalDate) : List<Map<String, Object>>",
                    "Materializes the requested date, then merges slots and bookings into available, blocked, booked, and operational states.",
                ],
            ],
        },
    }
    current_section = None
    pending_name = None
    tables_by_section: dict[str, dict[str, Table]] = {name: {} for name in sections}
    for block in list(iter_blocks(document)):
        if isinstance(block, Paragraph) and block.style.name.startswith("Heading 2"):
            current_section = block.text.strip() if block.text.strip() in sections else None
            pending_name = None
            continue
        if current_section is None:
            continue
        if isinstance(block, Paragraph):
            name = block.text.strip()
            pending_name = name if name in sections[current_section] else None
        elif isinstance(block, Table) and pending_name:
            set_table_rows(block, sections[current_section][pending_name])
            tables_by_section[current_section][pending_name] = block
            pending_name = None

    slot_generation_rows = [
        ["No", "Method", "Description"],
        [
            "01",
            "ensureDate(date : LocalDate) : int",
            "Materializes missing non-elapsed slots for every active field from the persisted opening, closing, and duration rules.",
        ],
        [
            "02",
            "rebuildRollingWindow() : int",
            "Rebuilds only unbooked available auto-generated rows and preserves bookings, Block exceptions, and manual audit rows.",
        ],
        [
            "03",
            "validateRuleChange(key : String, value : String) : void",
            "Validates opening/closing order, 30-minute duration steps, divisible operating windows, and the 1–90 day horizon.",
        ],
    ]
    search_anchor = tables_by_section["12. Search available fields by date/time"]["FieldOperationService"]
    insert_spec_after(document, search_anchor, "SlotGenerationService", slot_generation_rows)

    manage_anchor = tables_by_section["15. Manage unavailable slots"]["FieldOperationService"]
    policy_anchor = insert_spec_after(
        document,
        manage_anchor,
        "PromotionReportService",
        [
            ["No", "Method", "Description"],
            [
                "01",
                "updateSetting(key : String, value : String) : Map<String, Object>",
                "Lets Admin persist one automatic-slot rule after validating the combined rule set, then safely rebuilds the rolling window.",
            ],
        ],
    )
    insert_spec_after(document, policy_anchor, "SlotGenerationService", slot_generation_rows)

    calendar_anchor = tables_by_section["16. View field operation calendar"]["FieldOperationService"]
    insert_spec_after(document, calendar_anchor, "SlotGenerationService", slot_generation_rows[:2])


def update_sds_membership_rule_specs(document: Document) -> None:
    active = False
    pending_name = None
    membership_repository_table = None
    for block in list(iter_blocks(document)):
        if isinstance(block, Paragraph) and block.style.name.startswith("Heading 2"):
            active = block.text.strip() == "54. Manage membership rules"
            pending_name = None
            continue
        if not active:
            continue
        if isinstance(block, Paragraph):
            name = block.text.strip()
            pending_name = name if name in {
                "PromotionReportService",
                "MembershipLevelRepository",
            } else None
        elif isinstance(block, Table) and pending_name:
            if pending_name == "PromotionReportService":
                set_table_rows(
                    block,
                    [
                        ["No", "Method", "Description"],
                        [
                            "01",
                            "createMembershipLevel(request : ApiRequests.MembershipLevelUpsert) : Map<String, Object>",
                            "Validates and persists a new level, then recalculates every customer's existing membership assignment.",
                        ],
                        [
                            "02",
                            "updateMembershipLevel(id : Long, request : ApiRequests.MembershipLevelUpsert) : Map<String, Object>",
                            "Updates an existing rule and recalculates assignments without creating duplicate customer-membership rows.",
                        ],
                    ],
                )
            else:
                membership_repository_table = block
            pending_name = None

    if membership_repository_table is not None:
        insert_spec_after(
            document,
            membership_repository_table,
            "DomainSupportService",
            [
                ["No", "Method", "Description"],
                [
                    "01",
                    "refreshAllMembershipAssignments() : void",
                    "Re-evaluates all Customer accounts after an Admin changes a membership rule.",
                ],
                [
                    "02",
                    "updateMembershipProgress(customerId : Long, increment : int) : CustomerMembership",
                    "Updates the customer's single membership row in place and falls back to the first active tier when no threshold currently qualifies.",
                ],
                [
                    "03",
                    "attachDefaultMembership(customer : AppUser) : CustomerMembership",
                    "Returns the existing customer membership when present; otherwise inserts exactly one default assignment.",
                ],
            ],
        )


def update_sds_financial_specs(document: Document) -> None:
    sections = {
        "37. View checkout summary": {
            "BookingWorkflowService": [
                ["No", "Method", "Description"],
                [
                    "01",
                    "previewCheckout(request : ApiRequests.PromotionApply) : Map<String, Object>",
                    "Calculates the USD subtotal, caps one promotion at that subtotal, applies membership only when stacking is allowed, and derives deposit from the non-negative final total.",
                ],
            ],
            "DomainSupportService": [
                ["No", "Method", "Description"],
                [
                    "01",
                    "calculatePromotionDiscount(code, slot, baseAmount, services, customer) : BigDecimal",
                    "Validates eligibility and caps the promotion between zero and the eligible field-plus-service subtotal.",
                ],
                [
                    "02",
                    "calculateMembershipDiscount(customer, baseAmount, source) : BigDecimal",
                    "Applies the eligible online membership rate to a non-negative post-promotion balance.",
                ],
                [
                    "03",
                    "calculateDeposit(total : BigDecimal) : BigDecimal",
                    "Applies deposit.default_percent to the final non-negative USD total.",
                ],
            ],
        },
        "40. Capture/confirm online payment": {
            "PayPalCheckoutService": [
                ["No", "Method", "Description"],
                [
                    "01",
                    "captureOrder(bookingId, orderId, request) : Map<String, Object>",
                    "Uses a stable PayPal request identity, reconciles an ambiguous timeout by reading the order, validates exact USD capture, stores provider fee/net, updates gross paid balance, and regenerates the invoice once.",
                ],
            ],
        },
        "46. Process online refund": {
            "PayPalCheckoutService": [
                ["No", "Method", "Description"],
                [
                    "01",
                    "refundCapture(payment, refundAmount, requestId, existingRefundId) : RefundResult",
                    "Uses PayPal Payments v2 with a stable request id, polls an existing provider refund when present, and returns provider truth without inventing completion.",
                ],
            ],
            "PaymentWorkflowService": [
                ["No", "Method", "Description"],
                [
                    "01",
                    "updateRefundStatus(refundId, request) : Map<String, Object>",
                    "Maps only provider COMPLETED to completion, preserves booking paidAmount as gross collected, records refund separately, and synchronizes payment and invoice refund state.",
                ],
            ],
        },
        "47. Manage refund requests": {
            "PaymentWorkflowService": [
                ["No", "Method", "Description"],
                [
                    "01",
                    "createRefund(request : ApiRequests.RefundCreate) : Map<String, Object>",
                    "Validates Customer ownership and reserves only requested, approved, or processing capacity so completed partial refunds are not subtracted twice.",
                ],
                [
                    "02",
                    "updateRefundStatus(refundId, request) : Map<String, Object>",
                    "Validates operator transitions and completes either a PayPal provider refund or an explicit manual cash refund.",
                ],
                [
                    "03",
                    "refunds() : List<Map<String, Object>>",
                    "Returns newest-first requests with Staff/Admin scope or Customer ownership filtering.",
                ],
            ],
        },
        "51. Manage promotion campaigns": {
            "PromotionReportService": [
                ["No", "Method", "Description"],
                [
                    "01",
                    "createPromotion(request : ApiRequests.PromotionUpsert) : Map<String, Object>",
                    "Validates dates, positive discount, non-negative caps/minimum/usage limit, applicability, and stackable policy before creating a campaign.",
                ],
                [
                    "02",
                    "updatePromotion(id, request) : Map<String, Object>",
                    "Applies the same validation while updating campaign eligibility, dates, limits, status, and membership-stacking policy.",
                ],
            ],
        },
        "52. Apply promotion to booking": {
            "DomainSupportService": [
                ["No", "Method", "Description"],
                [
                    "01",
                    "calculatePromotionDiscount(code, slot, baseAmount, services, customer) : BigDecimal",
                    "Validates all campaign conditions and caps either percent or fixed discount at the non-negative subtotal.",
                ],
                [
                    "02",
                    "promotionAllowsMembershipStacking(code : String) : boolean",
                    "Treats stackable as permission to apply the Customer membership discount after the promotion.",
                ],
                [
                    "03",
                    "saveBookingPromotion(booking, code, discount) : void",
                    "Snapshots the applied code/discount and consumes campaign usage once when the booking is created.",
                ],
            ],
        },
    }
    current_section = None
    pending_name = None
    for block in list(iter_blocks(document)):
        if isinstance(block, Paragraph) and block.style.name.startswith("Heading 2"):
            current_section = block.text.strip() if block.text.strip() in sections else None
            pending_name = None
            continue
        if current_section is None:
            continue
        if isinstance(block, Paragraph):
            name = block.text.strip()
            pending_name = name if name in sections[current_section] else None
        elif isinstance(block, Table) and pending_name:
            set_table_rows(block, sections[current_section][pending_name])
            pending_name = None


def update_sds_walk_in_specs(document: Document) -> None:
    """Align final UC-25 with registered and first-time walk-in identities."""
    section_heading = "25. Create walk-in booking"
    target_tables = {
        "BookingWorkflowService": [
            ["No", "Method", "Description"],
            [
                "01",
                "createBooking(request : ApiRequests.BookingCreate) : Map<String, Object>",
                "Lets Venue Staff create a pending walk-in for either a matched Customer account or a first-time visitor name/phone snapshot.",
            ],
        ],
        "AccountService": [
            ["No", "Method", "Description"],
            [
                "01",
                "walkInCustomers() : List<Map<String, Object>>",
                "Provides the counter-safe active-Customer directory for optional phone/email matching; no match does not block a guest booking.",
            ],
        ],
        "DomainSupportService": [
            ["No", "Method", "Description"],
            [
                "01",
                "calculateMembershipDiscount(customer, baseAmount, source) : BigDecimal",
                "Returns zero membership discount for every walk-in, including a visitor without an account.",
            ],
            [
                "02",
                "bookingSummary(booking : Booking) : Map<String, Object>",
                "Returns registered contact data or the guest snapshot without exposing an unusable raw customer identifier.",
            ],
        ],
    }
    current_section = None
    pending_name = None
    drop_old_sql_tail = False
    service_tables: dict[str, Table] = {}
    for block in list(iter_blocks(document)):
        if isinstance(block, Paragraph) and block.style.name.startswith("Heading 2"):
            current_section = block.text.strip()
            pending_name = None
            drop_old_sql_tail = False
            continue
        if current_section != section_heading:
            continue
        if isinstance(block, Paragraph):
            text = block.text.strip()
            if drop_old_sql_tail:
                if text.startswith("Implementation note: The counter booking starts confirmed"):
                    block.text = (
                        "Implementation note: Exactly one identity is stored: customer_id, or for a first-time walk-in "
                        "guest_name and guest_phone (guest_email optional). Pay later stays pending. Deposit/full cash uses "
                        "PaymentWorkflowService.createWalkInAndCapture so booking and payment commit together; successful "
                        "deposit coverage confirms the booking. Guest walk-ins have no membership or PayPal flow."
                    )
                    drop_old_sql_tail = False
                else:
                    block._p.getparent().remove(block._p)
                continue
            pending_name = text if text in target_tables else None
            if text == "1. Create confirmed walk-in booking":
                block.text = "1. Create pending walk-in booking identity"
            elif text.startswith("INSERT INTO booking"):
                block.text = (
                    "INSERT INTO booking\n"
                    "  (customer_id, guest_name, guest_phone, guest_email, staff_id, slot_id,\n"
                    "   booking_code, status, booking_source, field_price_amount,\n"
                    "   service_total_amount, total_amount, deposit_amount, paid_amount,\n"
                    "   remaining_amount, created_at, updated_at)\n"
                    "VALUES\n"
                    "  (:customerId, :guestName, :guestPhone, :guestEmail, :currentStaffId, :slotId,\n"
                    "   :bookingCode, 'pending', 'walk_in', :fieldPrice, :serviceTotal,\n"
                    "   :totalAmount, :depositAmount, 0, :totalAmount, CURRENT_TIMESTAMP, CURRENT_TIMESTAMP);"
                )
                drop_old_sql_tail = True
        elif isinstance(block, Table) and pending_name:
            set_table_rows(block, target_tables[pending_name])
            service_tables[pending_name] = block
            pending_name = None

    booking_table = service_tables.get("BookingWorkflowService")
    if booking_table is None:
        raise RuntimeError("SDS UC-25 BookingWorkflowService specification was not found")
    insert_spec_after(
        document,
        booking_table,
        "PaymentWorkflowService",
        [
            ["No", "Method", "Description"],
            [
                "01",
                "createWalkInAndCapture(request : ApiRequests.WalkInCheckout) : Map<String, Object>",
                "Creates the walk-in and captures deposit/full cash in one transaction so a failed capture leaves no partial booking.",
            ],
            [
                "02",
                "capturePayment(request : ApiRequests.PaymentCapture) : Map<String, Object>",
                "Records later counter cash and confirms a pending booking once the required deposit is covered.",
            ],
        ],
    )


def update_sds() -> None:
    document = Document(SDS_INPUT)
    update_cover_metadata(document)
    replacements = [
        ("status, account_locked,", "status,"),
        ("'active', false, 0, false", "'active', 0, false"),
        ("status = 'locked', account_locked = true, lock_reason", "status = 'locked', lock_reason"),
        ("status = 'active', account_locked = false, lock_reason", "status = 'active', lock_reason"),
        ("Ask smart assistant for available fields", "Ask availability assistant"),
        (
            "Implementation note: Ranking and promotion eligibility are calculated in Java; no chatbot/RAG table is required.",
            "Implementation note: Availability, pricing, promotion eligibility, and membership discounts are deterministic Java rules over GoalZone data; no assistant or chatbot layer is in scope.",
        ),
        (
            "Implementation note: The API accepts a date and field type; time is represented by the returned slot start_time/end_time rather than a separate time request parameter.",
            "Implementation note: Missing availability is first materialized from the Admin opening, closing, duration, and horizon rules. The API then accepts a date and field type; time is represented by returned start_time/end_time values.",
        ),
        (
            "Implementation note: Blocking is rejected when an active booking already occupies the exact persisted slot.",
            "Implementation note: Block covers every overlapping generated slot and is rejected when any overlap has an active booking. A future exception may also be recorded before that date enters the rolling horizon.",
        ),
        (
            "Implementation note: The product target is PayPal Sandbox, not live-money certification. Mock mode exists for local automated tests.",
            "Implementation note: The product target is PayPal Sandbox, not live-money certification. PayPal receives the exact USD booking balance; processor fee/net are provider-returned merchant values, not customer tax. Mock mode exists only for automated tests.",
        ),
        (
            "Implementation note: Only provider COMPLETED capture is accepted as paid outside mock mode.",
            "Implementation note: Only provider COMPLETED capture is accepted as paid. Capture uses a stable PayPal-Request-Id and reads the order after an ambiguous timeout before changing the local gross-paid ledger.",
        ),
        (
            "Implementation note: SecurityConfig allows promotion mutations only for Admin. The scope supports one applied code per booking; stacking is out of scope even though the schema retains stackable.",
            "Implementation note: SecurityConfig allows promotion mutations only for Admin. One code is supported per booking; stackable explicitly controls whether the membership discount may also apply after the capped promotion.",
        ),
        (
            "Implementation note: Only PayPal provider COMPLETED becomes completed. Real buyer/refund confirmation remains an external sandbox smoke check, not live-money certification.",
            "Implementation note: Only PayPal provider COMPLETED becomes completed. paidAmount remains gross collected, refunds are recorded separately, and the stable provider request id makes retries safe.",
        ),
        (
            "Implementation note: This is a date-focused operations calendar, not a month-grid scheduling engine.",
            "Implementation note: The requested date is materialized before display. This remains a date-focused operations calendar, not a month-grid scheduling engine.",
        ),
        ("Flyway V11", "Flyway V20"),
        ("Flyway V15", "Flyway V20"),
        ("Flyway V16", "Flyway V20"),
        ("Flyway V17", "Flyway V20"),
        ("Flyway V18", "Flyway V20"),
    ]
    replace_text(document, replacements)
    replace_xml_attribute_text(
        document,
        [
            ("Flyway V1 through V11", "Flyway V1 through V20"),
            ("Flyway V1 through V15", "Flyway V1 through V20"),
            ("Flyway V1 through V16", "Flyway V1 through V20"),
            ("Flyway V1 through V17", "Flyway V1 through V20"),
            ("Flyway V1 through V18", "Flyway V1 through V20"),
            ("UC-64", "UC-63"),
        ],
    )
    update_sds_slot_generation_specs(document)
    update_sds_membership_rule_specs(document)
    update_sds_financial_specs(document)
    update_package_tables(document)
    update_database_table_descriptions(document)
    align_sds_retake_rules(document)

    replace_drawing_after(
        document,
        "1. Code Packages",
        "1. Code Packages",
        PACKAGE_IMAGE,
        max_height=6.8,
    )

    rebuild_sds_use_case_sections(document)
    update_sds_walk_in_specs(document)
    replace_all_sds_diagrams(document)

    replace_text(document, [("Flyway V20", "Flyway V21"), ("through V20", "through V21")])
    replace_xml_attribute_text(document, [("Flyway V20", "Flyway V21"), ("through V20", "through V21")])
    set_update_fields(document)
    normalize_document_font(document)
    OUTPUT.mkdir(parents=True, exist_ok=True)
    document.save(SDS_OUTPUT)


if __name__ == "__main__":
    update_rds()
    update_sds()
    print(RDS_OUTPUT)
    print(SDS_OUTPUT)
