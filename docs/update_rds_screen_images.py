#!/usr/bin/env python3
"""Replace aligned RDS screen-design captures without rebuilding other sections."""

from __future__ import annotations

import os
from pathlib import Path
import zipfile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
RDS = ROOT / "output/doc/GoalZone_RDS_Retake_Aligned_2026-08-01.docx"
SCREENSHOTS = ROOT / "output/playwright/rds-final"
SCREEN_STATES = ROOT / "output/playwright/rds-screen-states"
TMP = ROOT / "tmp/docs"

STATE_CAPTION_PREFIX = "Screen state - "
ADDITIONAL_SCREEN_STATES = {
    "e. Profile / Edit Profile": [
        ("Change password form", "customer-change-password.jpg"),
    ],
    "f. Customer Management": [
        ("Customer activity detail", "admin-customer-activity-modal.jpg"),
        ("Customer lock confirmation", "admin-customer-lock-modal.jpg"),
    ],
    "e. Field Pricing Management": [
        ("Pricing rule list and editor", "admin-field-pricing.jpg"),
    ],
    "g. Extra Service Management": [
        ("Service list and editor", "admin-services-editor.jpg"),
    ],
    "h. Unavailable Slot Management": [
        ("Daily field schedule", "staff-field-schedule-list.jpg"),
        ("Block slot form", "staff-field-block-form.jpg"),
    ],
    "i. Issue Report / Issue Management": [
        ("Customer issue report form", "customer-issue-report-form.jpg"),
        ("Issue details and submit action", "customer-issue-report-details.jpg"),
        ("Staff issue handling", "staff-issue-management.jpg"),
    ],
    "a. Online Booking": [
        ("Slot services and checkout summary", "customer-services-checkout.jpg"),
        ("Online payment choice", "customer-payment-choice.jpg"),
    ],
    "b. Walk-in Booking": [
        ("Selected slot and services", "walkin-slot-services.jpg"),
        ("Counter payment summary", "walkin-payment-summary.jpg"),
    ],
    "c. Booking Detail": [
        ("Booking list and selected detail", "customer-booking-list-detail.jpg"),
        ("Invoice and transaction history", "customer-booking-invoice.jpg"),
    ],
    "e. Booking Operations": [
        ("Booking queue and selected operation", "staff-booking-queue-operation.jpg"),
        ("Schedule and counter payment actions", "staff-booking-schedule-payment.jpg"),
        ("Invoice and transaction disclosure", "staff-booking-invoice-history.jpg"),
    ],
    "f. Reschedule Booking": [
        ("Reschedule controls and available slot selection", "customer-change-cancellation.jpg"),
    ],
    "g. Cancellation Preview": [
        ("Policy-derived cancellation terms", "customer-change-cancellation.jpg"),
    ],
    "a. Checkout Summary": [
        ("Service totals and payment choices", "customer-services-checkout.jpg"),
    ],
    "b. Payment Sandbox": [
        ("Customer online payment selection", "customer-payment-choice.jpg"),
    ],
    "f. Policy Management": [
        ("Booking, cancellation, and notification policies", "admin-booking-policies.jpg"),
        ("Automatic slot generation policies", "admin-slot-generation-policies.jpg"),
    ],
    "b. Promotion Management": [
        ("Promotion editor", "admin-promotion-edit-modal.jpg"),
    ],
    "e. Membership Rule Management": [
        ("Membership tier editor", "admin-membership-tier-edit-modal.jpg"),
    ],
    "f. Notification List": [
        ("Customer notification list", "customer-notifications.jpg"),
    ],
    "g. Revenue Report": [
        ("Revenue range and summary", "admin-overview-revenue.jpg"),
    ],
    "h. Booking Report": [
        ("Booking trends and field utilization", "admin-overview-demand-customers.jpg"),
    ],
    "i. Customer Account State Report": [
        ("Returning customers and membership distribution", "admin-overview-demand-customers.jpg"),
    ],
}

TEXT_REPLACEMENTS = {
    "UC-07/UC-08, UC-10": "UC-07, UC-08, UC-10",
    "UC-16/UC-17, UC-18": "UC-17, UC-18",
    "Related use cases: UC-16, UC-17, UC-18": "Related use cases: UC-17, UC-18",
    "Detailed as-built SQL is specified under UC-16, UC-17, UC-18 in the SDS Database Queries sections.":
        "Detailed as-built SQL is specified under UC-17, UC-18 in the SDS Database Queries sections.",
    "UC-20, UC-20, UC-21": "UC-20, UC-21",
    "UC-25, UC-36, UC-37": "UC-25, UC-35, UC-37",
    "UC-26, UC-36, UC-24, UC-31, UC-32, UC-33": "UC-26, UC-31, UC-32, UC-33, UC-37, UC-40",
    "UC-28, UC-36, UC-24, UC-31, UC-32, UC-33, UC-29/UC-30": "UC-28, UC-29, UC-30, UC-31, UC-32, UC-33, UC-37",
    "UC-29, UC-29/UC-30": "UC-29",
    "UC-30, UC-30": "UC-30",
    "UC-36, UC-36, UC-38": "UC-36, UC-38",
    "UC-40, UC-40": "UC-40",
    "Customer Account state Report": "Customer Account State Report",
    "e. Booking Calendar": "e. Booking Operations",
    "f. Booking and Refund Policies": "f. Policy Management",
    "Staff daily operations queue/calendar with selected-booking controls.":
        "Staff booking queue with selected-booking lifecycle, schedule, payment, invoice, and service controls.",
    "Staff blocks/unblocks field time with a reason and reviews daily field operations.":
        "Venue Staff reviews daily field operations and blocks or unblocks operational slot exceptions; Admin configures ordinary slot generation separately.",
    "Customer/Staff reports a field, booking, or service issue; Staff resolves it and notifies the reporter.":
        "Customer or Staff reports a field, booking, or service issue; Venue Staff resolves or rejects it and notifies the reporter; Admin reviews the history in read-only mode.",
    "Role/ownership-protected view of booking, billing, and allowed lifecycle actions.":
        "Customer views an owned booking; Venue Staff performs eligible lifecycle actions; Admin reviews booking and billing data in read-only audit mode.",
    "Shows real persisted transaction history subject to booking ownership/operator access.":
        "Shows persisted transaction history subject to ownership or operational access; Admin access is read-only audit.",
    "Displays the generated booking invoice and current settlement state.":
        "Displays the generated booking invoice and current settlement state; Admin access is read-only audit.",
    "Customer submits an eligible request; Staff reviews it and completes PayPal or cash refund workflows.":
        "Customer submits an eligible request; Venue Staff reviews and completes PayPal or cash refund workflows; Admin reviews the audit trail without changing it.",
    "Admin maintains deposit, payment-timeout, cancellation/refund, and reminder settings.":
        "Admin maintains slot-generation, deposit, payment-timeout, cancellation/refund, and reminder settings.",
    "Related use cases: UC-43, UC-44": "Related use cases: UC-16, UC-43, UC-44",
    "Detailed as-built SQL is specified under UC-43, UC-44 in the SDS Database Queries sections.":
        "Detailed as-built SQL is specified under UC-16, UC-43, UC-44 in the SDS Database Queries sections.",
}

SCREEN_DESCRIPTION_UPDATES = {
    "Unavailable Slot Management": (
        "Venue Staff reviews the daily field schedule and blocks or unblocks operational slot exceptions. "
        "Admin slot-generation rules are configured separately in the policy workspace."
    ),
    "Issue Report / Issue Management": (
        "Customer or Staff reports a field, booking, or service issue; Venue Staff resolves or rejects it; "
        "Admin has read-only audit access."
    ),
    "Walk-in Booking": (
        "Venue Staff creates a walk-in booking for either a Customer matched by phone/email or a first-time "
        "visitor identified by name and phone, then records deposit/full cash or leaves payment pending."
    ),
    "Booking Detail": (
        "Customer views an owned booking, Venue Staff performs eligible lifecycle actions, and Admin reviews "
        "booking and billing data in read-only audit mode."
    ),
    "Booking Operations": (
        "Venue Staff uses the booking queue and selected-booking controls for reschedule, cancellation, check-in, "
        "completion, no-show, services, cash payment, invoice, and transaction history."
    ),
    "Payment History": (
        "Customer views owned transactions; Venue Staff and Admin review role-authorized payment records, with "
        "Admin access remaining read-only."
    ),
    "Invoice Detail": (
        "Displays the generated invoice and settlement state to the booking owner or Venue Staff; Admin access is "
        "read-only audit."
    ),
    "Refund Management": (
        "Customer submits an eligible refund request, Venue Staff reviews and processes it, and Admin inspects "
        "the refund audit trail without changing it."
    ),
    "Policy Management": (
        "Admin policy workspace for deposit, cancellation/refund, payment timeout, reminder, and automatic slot "
        "opening/closing/duration/horizon rules."
    ),
}

def all_paragraphs(document: Document):
    yield from document.paragraphs
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def normalize_document_font(document: Document, font_name: str = "Arial") -> None:
    for style in document.styles:
        if getattr(style, "font", None) is None:
            continue
        style.font.name = font_name
        style_properties = style.element.get_or_add_rPr()
        style_fonts = style_properties.rFonts
        if style_fonts is None:
            style_fonts = OxmlElement("w:rFonts")
            style_properties.insert(0, style_fonts)
        for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
            style_fonts.set(qn(f"w:{attribute}"), font_name)

    for part in document.part.package.parts:
        element = getattr(part, "element", None)
        if element is None:
            continue
        for run in element.xpath(".//w:r"):
            run_properties = run.get_or_add_rPr()
            run_fonts = run_properties.rFonts
            if run_fonts is None:
                run_fonts = OxmlElement("w:rFonts")
                run_properties.insert(0, run_fonts)
            for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
                run_fonts.set(qn(f"w:{attribute}"), font_name)


def remove_generated_screen_states(document: Document) -> None:
    body = document.element.body
    for paragraph in list(document.paragraphs):
        if not paragraph.text.strip().startswith(STATE_CAPTION_PREFIX):
            continue
        caption_element = paragraph._p
        image_element = caption_element.getnext()
        if image_element is not None and image_element.tag.endswith("}p") and image_element.xpath(".//w:drawing"):
            body.remove(image_element)
        body.remove(caption_element)


def insert_additional_screen_states(document: Document) -> None:
    missing = [
        str(SCREEN_STATES / filename)
        for states in ADDITIONAL_SCREEN_STATES.values()
        for _, filename in states
        if not (SCREEN_STATES / filename).is_file()
    ]
    if missing:
        raise FileNotFoundError("Missing additional screen states: " + ", ".join(missing))

    screen_headings = {paragraph.text.strip(): paragraph for paragraph in document.paragraphs}
    absent_headings = sorted(set(ADDITIONAL_SCREEN_STATES) - set(screen_headings))
    if absent_headings:
        raise RuntimeError("RDS screen headings changed unexpectedly: " + ", ".join(absent_headings))

    for heading, states in ADDITIONAL_SCREEN_STATES.items():
        cursor = screen_headings[heading]._p
        element = cursor.getnext()
        while element is not None:
            if element.tag.endswith("}p"):
                paragraph_text = "".join(node.text or "" for node in element.xpath(".//w:t")).strip()
                if paragraph_text == "UI Design":
                    image_element = element.getnext()
                    while image_element is not None and not image_element.xpath(".//w:drawing"):
                        image_element = image_element.getnext()
                    if image_element is None:
                        raise RuntimeError(f"UI image missing after {heading}")
                    cursor = image_element
                    break
            element = element.getnext()
        else:
            raise RuntimeError(f"UI Design marker missing after {heading}")

        for label, filename in states:
            caption = document.add_paragraph()
            caption.paragraph_format.keep_with_next = True
            caption.paragraph_format.space_before = Pt(4)
            caption.paragraph_format.space_after = Pt(2)
            caption_run = caption.add_run(f"{STATE_CAPTION_PREFIX}{label}")
            caption_run.bold = True
            caption_run.italic = True
            caption_run.font.size = Pt(9)

            picture = document.add_paragraph()
            picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
            picture.paragraph_format.space_after = Pt(4)
            picture.add_run().add_picture(str(SCREEN_STATES / filename), width=Inches(6.5))

            cursor.addnext(caption._p)
            caption._p.addnext(picture._p)
            cursor = picture._p


def clean_screen_references(document: Document) -> None:
    for paragraph in all_paragraphs(document):
        updated = paragraph.text
        for old, new in TEXT_REPLACEMENTS.items():
            updated = updated.replace(old, new)
        if updated != paragraph.text:
            paragraph.text = updated

    for table in document.tables:
        if not table.rows:
            continue
        header = [cell.text.strip() for cell in table.rows[0].cells]
        if header[:4] == ["#", "Feature", "Screen", "Description"]:
            for row in table.rows[1:]:
                screen_name = row.cells[2].text.strip()
                if screen_name == "Booking Calendar":
                    screen_name = "Booking Operations"
                    row.cells[2].text = screen_name
                elif screen_name == "Booking and Refund Policies":
                    screen_name = "Policy Management"
                    row.cells[2].text = screen_name
                if screen_name in SCREEN_DESCRIPTION_UPDATES:
                    row.cells[3].text = SCREEN_DESCRIPTION_UPDATES[screen_name]
        elif header[:5] == ["Screen", "Guest", "Customer", "Staff", "Admin"]:
            for row in table.rows[1:]:
                if row.cells[0].text.strip() == "Booking Calendar":
                    row.cells[0].text = "Booking Operations"
                elif row.cells[0].text.strip() == "Booking and Refund Policies":
                    row.cells[0].text = "Policy Management"

    for table in document.tables:
        if not table.rows:
            continue
        header = [cell.text.strip() for cell in table.rows[0].cells]
        rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
        if header[:3] != ["Field Name", "Field Type", "Description"] or not any(
            row.startswith(("Customer |", "Customer type |")) for row in rows
        ):
            continue
        values = [
            ("Customer type", "Segmented buttons", "Choose a registered Customer or a first-time visitor."),
            ("Phone or email", "Search box", "Optionally matches an active registered Customer at the counter."),
            ("Guest contact", "Text inputs", "First-time visitor name and phone are required; email is optional."),
            ("Field/date/slot", "Selection", "Selects an available non-past slot."),
            ("Services/promotion", "Selection", "Optional validated add-ons and promotion."),
            ("Payment option", "Choice", "Deposit/full cash confirms the booking; Pay later keeps it pending."),
            ("Create booking", "Button", "Stores exactly one customer identity and creates the walk-in safely."),
        ]
        while len(table.rows) < len(values) + 1:
            table.add_row()
        for row, row_values in zip(table.rows[1:], values):
            for cell, value in zip(row.cells, row_values):
                cell.text = value
        break

    for table in document.tables:
        if not table.rows:
            continue
        header = [cell.text.strip() for cell in table.rows[0].cells]
        rows = [tuple(cell.text.strip() for cell in row.cells) for row in table.rows]
        if header[:3] != ["Table", "CRUD", "Description"] or not any(
            "app_user" in row[0] and "booking" in row[0] and "slot" in row[0] for row in rows[1:]
        ):
            continue
        for row in table.rows[1:]:
            entity_names = row.cells[0].text.strip()
            if "app_user" in entity_names and "booking" in entity_names and "slot" in entity_names:
                row.cells[1].text = "R/C/U"
                row.cells[2].text = (
                    "Validates Staff and an optional matched Customer, records either registered or guest "
                    "identity, and creates a pending or confirmed walk-in booking."
                )
            elif "payment" in entity_names and "invoice" in entity_names:
                row.cells[2].text = "Stores cash when collected and maintains invoice and billing state."
        break


def replace_media(package_source: Path, destination: Path) -> None:
    replacement_entries = {
        f"word/media/image{number}.png": SCREENSHOTS / f"image{number}.png"
        for number in range(10, 45)
    }
    missing = [str(path) for path in replacement_entries.values() if not path.is_file()]
    if missing:
        raise FileNotFoundError("Missing screen captures: " + ", ".join(missing))

    package_tmp = TMP / f"{destination.stem}.package.tmp.docx"
    with zipfile.ZipFile(package_source, "r") as source, zipfile.ZipFile(package_tmp, "w") as target:
        package_names = set(source.namelist())
        absent_parts = sorted(set(replacement_entries) - package_names)
        if absent_parts:
            raise RuntimeError("RDS media parts changed unexpectedly: " + ", ".join(absent_parts))
        for info in source.infolist():
            replacement = replacement_entries.get(info.filename)
            payload = replacement.read_bytes() if replacement else source.read(info.filename)
            target.writestr(info, payload)
    os.replace(package_tmp, destination)


def main() -> None:
    TMP.mkdir(parents=True, exist_ok=True)
    document = Document(RDS)
    remove_generated_screen_states(document)
    clean_screen_references(document)
    insert_additional_screen_states(document)
    normalize_document_font(document)
    xml_tmp = TMP / f"{RDS.stem}.xml-edited.docx"
    document.save(xml_tmp)
    replace_media(xml_tmp, RDS)

    verified = Document(RDS)
    expected_inline_shapes = 42 + sum(len(states) for states in ADDITIONAL_SCREEN_STATES.values())
    if len(verified.inline_shapes) != expected_inline_shapes:
        raise RuntimeError(
            f"Expected {expected_inline_shapes} inline drawings after update, found {len(verified.inline_shapes)}"
        )
    print(f"Updated {RDS} with the final UI screenshot set and aligned screen references.")


if __name__ == "__main__":
    main()
